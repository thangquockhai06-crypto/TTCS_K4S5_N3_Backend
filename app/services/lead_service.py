import io
import re
import csv
import uuid
from datetime import datetime, timedelta
from typing import List, Optional, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import desc, func, or_
from fastapi import HTTPException, status
from fastapi.responses import StreamingResponse

from app.models.lead import Lead, LeadRoutingRule
from app.models.user import User
from app.models.customer import Customer
from app.models.contact import Contact
from app.models.deal import Deal
from app.models.campaign import Campaign
from app.models.audit_log import AuditLog
from app.repositories.lead_repository import LeadRepository
from app.repositories.campaign_repository import CampaignRepository
from app.schemas.lead import (
    CreateLeadDTO,
    UpdateLeadDTO,
    PublicLeadSubmissionDTO,
    LeadDTO,
    MergeLeadsDTO,
    ConvertLeadDTO,
    RejectLeadDTO,
)

# Email domain free (không phải email doanh nghiệp)
FREE_EMAIL_DOMAINS = {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "aol.com", "icloud.com"}
EXECUTIVE_TITLES = {
    "giám đốc", "director", "ceo", "cfo", "cto", "vp", "vice president",
    "phó giám đốc", "trưởng phòng", "head of", "quản lý", "manager", "chủ tịch", "founder", "co-founder"
}


class LeadService:
    """
    Tầng Service xử lý toàn bộ nghiệp vụ Quản lý Đầu mối bán hàng (EP-04: S4-01 -> S4-09).
    """

    @staticmethod
    def calculate_score(lead_data: Dict[str, Any]) -> Tuple[int, str]:
        """
        S4-05: Thuật toán chấm điểm Lead động dựa trên các tiêu chí nghiệp vụ thực tế:
        - Email doanh nghiệp (không dùng gmail/yahoo): +15 điểm
        - Số điện thoại hợp lệ: +10 điểm
        - Có tên công ty: +15 điểm
        - Chức danh cấp quản lý/ra quyết định (Director, Manager, CEO): +25 điểm
        - Quy mô công ty >= 50 nhân sự: +15 điểm
        - Có ngân sách dự kiến > 50 triệu: +20 điểm
        Phân loại: Hot (>= 70), Warm (40 - 69), Cold (< 40).
        """
        score = 0
        email = (lead_data.get("email") or "").lower().strip()
        if email and "@" in email:
            domain = email.split("@")[-1]
            if domain not in FREE_EMAIL_DOMAINS:
                score += 15
            else:
                score += 5

        phone = lead_data.get("phone")
        if phone and len(str(phone).strip()) >= 9:
            score += 10

        company = lead_data.get("company")
        if company and len(str(company).strip()) >= 3:
            score += 15

        job_title = (lead_data.get("job_title") or lead_data.get("jobTitle") or "").lower()
        if any(title in job_title for title in EXECUTIVE_TITLES):
            score += 25
        elif job_title:
            score += 10

        company_size = str(lead_data.get("company_size") or lead_data.get("companySize") or "")
        if any(s in company_size for s in ["51-200", "201-500", "500+"]):
            score += 15
        elif company_size:
            score += 5

        annual_budget = float(lead_data.get("annual_budget") or lead_data.get("annualBudget") or 0.0)
        if annual_budget >= 50000000.0:
            score += 20
        elif annual_budget > 0:
            score += 10

        # Giới hạn điểm 0 - 100
        score = max(0, min(100, score))

        if score >= 70:
            grade = "Hot"
        elif score >= 40:
            grade = "Warm"
        else:
            grade = "Cold"

        return score, grade

    @staticmethod
    def auto_route_lead(db: Session, lead: Lead) -> Optional[User]:
        """
        S4-06: Tự động phân bổ Lead theo thứ tự ưu tiên:
        1. Vùng miền (Region)
        2. Ngành nghề (Industry)
        3. Phân bổ xoay vòng (Round-Robin) cho Sales Reps đang hoạt động.
        """
        rules = LeadRepository.get_routing_rules(db)

        # 1. Khớp theo Vùng miền (Region)
        if lead.region:
            for rule in rules:
                if rule.rule_type == "REGION" and rule.criterion_value:
                    if rule.criterion_value.lower() in lead.region.lower():
                        if rule.target_user_id:
                            target_u = db.query(User).filter(User.id == rule.target_user_id, User.status == "active").first()
                            if target_u:
                                return target_u

        # 2. Khớp theo Ngành nghề (Industry)
        if lead.industry:
            for rule in rules:
                if rule.rule_type == "INDUSTRY" and rule.criterion_value:
                    if rule.criterion_value.lower() in lead.industry.lower():
                        if rule.target_user_id:
                            target_u = db.query(User).filter(User.id == rule.target_user_id, User.status == "active").first()
                            if target_u:
                                return target_u

        # 3. Phân bổ Round-Robin
        return LeadRepository.get_next_round_robin_sales_user(db)

    @classmethod
    def create_lead(cls, db: Session, dto: CreateLeadDTO, user: Optional[User] = None) -> Lead:
        """
        S4-02: Tạo thủ công Đầu mối mới có chấm điểm và phân bổ tự động.
        """
        score, grade = cls.calculate_score(dto.model_dump())

        assigned_u = None
        assigned_user_id = dto.assignedUserId
        if assigned_user_id:
            assigned_u = db.query(User).filter(User.id == assigned_user_id).first()

        now = datetime.utcnow()
        new_lead = Lead(
            id=str(uuid.uuid4()),
            full_name=dto.fullName.strip(),
            email=dto.email.strip().lower(),
            phone=dto.phone.strip() if dto.phone else None,
            company=dto.company.strip() if dto.company else None,
            job_title=dto.jobTitle,
            source=dto.source or "Manual",
            campaign_id=dto.campaignId,
            region=dto.region or "Miền Bắc",
            industry=dto.industry or "Công nghệ",
            company_size=dto.companySize or "11-50",
            annual_budget=dto.annualBudget or 0.0,
            score=score,
            score_grade=grade,
            status="ASSIGNED" if assigned_u else "UNASSIGNED",
            assigned_user_id=assigned_u.id if assigned_u else None,
            assigned_at=now if assigned_u else None,
            sla_deadline=now + timedelta(hours=24),  # SLA mặc định 24 giờ phản hồi
            sla_overdue=False,
            notes=dto.notes,
            created_at=now,
            updated_at=now,
        )

        if not assigned_u:
            # Tự động gán nếu chưa chọn người phụ trách
            auto_u = cls.auto_route_lead(db, new_lead)
            if auto_u:
                new_lead.assigned_user_id = auto_u.id
                new_lead.assigned_at = now
                new_lead.status = "ASSIGNED"

        return LeadRepository.create(db, new_lead)

    @classmethod
    def submit_public_lead(cls, db: Session, dto: PublicLeadSubmissionDTO) -> Lead:
        """
        S4-01: Tiếp nhận Lead từ Form nhúng công khai (Website Lead Capture):
        - Không để lộ secret key hay yêu cầu đăng nhập.
        - Tự động gắn Nguồn 'Web Form' và liên kết Campaign nếu có mã.
        - Tự động tính điểm và định tuyến gán nhân viên phụ trách.
        - Khởi tạo SLA 24h.
        """
        campaign_id = None
        if dto.campaignCode:
            c = CampaignRepository.get_by_code(db, dto.campaignCode.strip().upper())
            if c:
                campaign_id = c.id

        score, grade = cls.calculate_score({
            "email": dto.email,
            "phone": dto.phone,
            "company": dto.company,
        })

        now = datetime.utcnow()
        new_lead = Lead(
            id=str(uuid.uuid4()),
            full_name=dto.fullName.strip(),
            email=dto.email.strip().lower(),
            phone=dto.phone.strip() if dto.phone else None,
            company=dto.company.strip() if dto.company else None,
            source="Web Form",
            campaign_id=campaign_id,
            region="Miền Bắc",
            industry="Công nghệ",
            score=score,
            score_grade=grade,
            status="UNASSIGNED",
            sla_deadline=now + timedelta(hours=24),
            sla_overdue=False,
            notes=dto.message,
            created_at=now,
            updated_at=now,
        )

        # Định tuyến tự động
        auto_u = cls.auto_route_lead(db, new_lead)
        if auto_u:
            new_lead.assigned_user_id = auto_u.id
            new_lead.assigned_at = now
            new_lead.status = "ASSIGNED"

        return LeadRepository.create(db, new_lead)

    @staticmethod
    def map_to_dto(lead: Lead) -> LeadDTO:
        now = datetime.utcnow()
        remaining_hours = None
        if lead.sla_deadline:
            delta = lead.sla_deadline - now
            remaining_hours = round(delta.total_seconds() / 3600.0, 1)

        assigned_name = lead.assigned_user.full_name if lead.assigned_user else None
        campaign_name = lead.campaign.name if lead.campaign else None

        return LeadDTO(
            id=lead.id,
            fullName=lead.full_name,
            email=lead.email,
            phone=lead.phone,
            company=lead.company,
            jobTitle=lead.job_title,
            source=lead.source,
            campaignId=lead.campaign_id,
            campaignName=campaign_name,
            region=lead.region,
            industry=lead.industry,
            companySize=lead.company_size,
            annualBudget=float(lead.annual_budget or 0.0),
            score=lead.score,
            scoreGrade=lead.score_grade,
            status=lead.status,
            assignedUserId=lead.assigned_user_id,
            assignedUserName=assigned_name,
            assignedAt=lead.assigned_at.isoformat() if lead.assigned_at else None,
            rejectionReason=lead.rejection_reason,
            slaDeadline=lead.sla_deadline.isoformat() if lead.sla_deadline else None,
            slaOverdue=lead.sla_overdue,
            slaRemainingHours=remaining_hours,
            convertedCustomerId=lead.converted_customer_id,
            convertedDealId=lead.converted_deal_id,
            convertedAt=lead.converted_at.isoformat() if lead.converted_at else None,
            notes=lead.notes,
            createdAt=lead.created_at.isoformat() if lead.created_at else None,
            updatedAt=lead.updated_at.isoformat() if lead.updated_at else None,
        )

    @classmethod
    def get_leads(
        cls,
        db: Session,
        user: Optional[User] = None,
        search: Optional[str] = None,
        status: Optional[str] = None,
        score_grade: Optional[str] = None,
        source: Optional[str] = None,
        campaign_id: Optional[str] = None,
        sla_overdue: Optional[bool] = None,
        skip: int = 0,
        limit: int = 50,
    ) -> Tuple[List[LeadDTO], int]:
        leads = LeadRepository.get_all(
            db=db,
            user=user,
            search=search,
            status=status,
            score_grade=score_grade,
            source=source,
            campaign_id=campaign_id,
            sla_overdue=sla_overdue,
            skip=skip,
            limit=limit,
        )
        total = LeadRepository.count(
            db=db,
            user=user,
            search=search,
            status=status,
            score_grade=score_grade,
            source=source,
            campaign_id=campaign_id,
            sla_overdue=sla_overdue,
        )
        return [cls.map_to_dto(l) for l in leads], total

    @classmethod
    def get_lead_by_id(cls, db: Session, lead_id: str) -> LeadDTO:
        lead = LeadRepository.get_by_id(db, lead_id)
        if not lead:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy đầu mối với mã ID '{lead_id}'.",
            )
        return cls.map_to_dto(lead)

    @classmethod
    def update_lead(cls, db: Session, lead_id: str, dto: UpdateLeadDTO) -> LeadDTO:
        lead = LeadRepository.get_by_id(db, lead_id)
        if not lead:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy đầu mối với mã ID '{lead_id}'.",
            )

        update_dict = dto.model_dump(exclude_unset=True)
        if "fullName" in update_dict and update_dict["fullName"]:
            lead.full_name = update_dict["fullName"].strip()
        if "email" in update_dict and update_dict["email"]:
            lead.email = update_dict["email"].strip().lower()
        if "phone" in update_dict:
            lead.phone = update_dict["phone"]
        if "company" in update_dict:
            lead.company = update_dict["company"]
        if "jobTitle" in update_dict:
            lead.job_title = update_dict["jobTitle"]
        if "source" in update_dict:
            lead.source = update_dict["source"]
        if "campaignId" in update_dict:
            lead.campaign_id = update_dict["campaignId"]
        if "region" in update_dict:
            lead.region = update_dict["region"]
        if "industry" in update_dict:
            lead.industry = update_dict["industry"]
        if "companySize" in update_dict:
            lead.company_size = update_dict["companySize"]
        if "annualBudget" in update_dict:
            lead.annual_budget = update_dict["annualBudget"]
        if "status" in update_dict and update_dict["status"]:
            lead.status = update_dict["status"].upper()
        if "assignedUserId" in update_dict:
            lead.assigned_user_id = update_dict["assignedUserId"]
            if lead.assigned_user_id and not lead.assigned_at:
                lead.assigned_at = datetime.utcnow()
        if "notes" in update_dict:
            lead.notes = update_dict["notes"]

        # Recalculate score
        lead.score, lead.score_grade = cls.calculate_score({
            "email": lead.email,
            "phone": lead.phone,
            "company": lead.company,
            "job_title": lead.job_title,
            "company_size": lead.company_size,
            "annual_budget": lead.annual_budget,
        })

        saved = LeadRepository.update(db, lead)
        return cls.map_to_dto(saved)

    @classmethod
    def accept_lead(cls, db: Session, lead_id: str, user: User) -> LeadDTO:
        """
        S4-07: Tiếp nhận Lead:
        Chuyển trạng thái UNASSIGNED -> ASSIGNED -> CONTACTED và đánh dấu SLA phản hồi.
        """
        lead = LeadRepository.get_by_id(db, lead_id)
        if not lead:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy đầu mối với mã ID '{lead_id}'.",
            )
        now = datetime.utcnow()
        lead.assigned_user_id = user.id
        lead.assigned_at = now
        lead.status = "CONTACTED"
        lead.sla_responded_at = now
        saved = LeadRepository.update(db, lead)
        return cls.map_to_dto(saved)

    @classmethod
    def reject_lead(cls, db: Session, lead_id: str, dto: RejectLeadDTO, user: User) -> LeadDTO:
        """
        S4-07: Từ chối Lead với lý do bắt buộc.
        """
        lead = LeadRepository.get_by_id(db, lead_id)
        if not lead:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy đầu mối với mã ID '{lead_id}'.",
            )
        lead.status = "REJECTED"
        lead.rejection_reason = dto.rejectionReason.strip()
        saved = LeadRepository.update(db, lead)
        return cls.map_to_dto(saved)

    @classmethod
    def convert_lead(cls, db: Session, lead_id: str, dto: ConvertLeadDTO, user: User) -> Dict[str, Any]:
        """
        S4-08: Chuyển đổi Lead thành Customer, Contact và Deal trong một transaction an toàn:
        - Tái sử dụng Customer hiện có nếu truyền existingCustomerId.
        - Tạo mới Customer nếu chưa có.
        - Tự động tạo bản ghi Contact tương ứng.
        - Tạo Deal cơ hội bán hàng nếu createDeal = True.
        - Đánh dấu lead là CONVERTED chỉ khi mọi bước thành công.
        """
        lead = LeadRepository.get_by_id(db, lead_id)
        if not lead:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy đầu mối với mã ID '{lead_id}'.",
            )

        if lead.status == "CONVERTED":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Đầu mối này đã được chuyển đổi trước đó. Không thể chuyển đổi lại.",
            )

        now = datetime.utcnow()

        try:
            # 1. Khách hàng Doanh nghiệp (Customer)
            customer = None
            if dto.existingCustomerId:
                customer = db.query(Customer).filter(Customer.id == dto.existingCustomerId).first()
            if not customer:
                customer_name = dto.customerName or lead.company or lead.full_name
                customer = Customer(
                    id=str(uuid.uuid4()),
                    full_name=customer_name,
                    email=lead.email,
                    phone=lead.phone or "0912345678",
                    company=lead.company or customer_name,
                    status="prospect",
                    health_score=85,
                    assigned_user_id=lead.assigned_user_id or user.id,
                    tax_code=dto.taxCode,
                    industry=lead.industry or "Công nghệ",
                    tier="Standard",
                    created_at=now,
                    updated_at=now,
                )
                db.add(customer)
                db.flush()

            # 2. Người liên hệ chính (Contact)
            contact = Contact(
                id=str(uuid.uuid4()),
                customer_id=customer.id,
                full_name=lead.full_name,
                email=lead.email,
                phone=lead.phone,
                position=lead.job_title or "Đại diện đầu mối",
                role="Decider",
                is_primary=True,
                is_active=True,
                notes=lead.notes,
                created_at=now,
                updated_at=now,
            )
            db.add(contact)
            db.flush()

            # 3. Cơ hội bán hàng (Deal) nếu được yêu cầu
            deal = None
            if dto.createDeal:
                deal_title = dto.dealTitle or f"Cơ hội từ {lead.full_name}"
                deal_val = float(dto.dealValue or lead.annual_budget or 50000000.0)
                from datetime import date
                close_d = date.fromisoformat(dto.expectedCloseDate) if dto.expectedCloseDate else None

                deal = Deal(
                    id=str(uuid.uuid4()),
                    title=deal_title,
                    value=deal_val,
                    stage="lead",
                    probability=25,
                    customer_id=customer.id,
                    owner_id=lead.assigned_user_id or user.id,
                    expected_close_date=close_d,
                    campaign_id=lead.campaign_id,
                    created_at=now,
                    updated_at=now,
                )
                db.add(deal)
                db.flush()

            # 4. Cập nhật Lead thành CONVERTED
            lead.status = "CONVERTED"
            lead.converted_customer_id = customer.id
            lead.converted_contact_id = contact.id
            lead.converted_deal_id = deal.id if deal else None
            lead.converted_at = now
            db.commit()

            # Ghi nhật ký kiểm toán (AuditLog)
            audit = AuditLog(
                id=str(uuid.uuid4()),
                user_id=user.id,
                action="LEAD_CONVERTED",
                performed_by=user.full_name,
                user_name=user.full_name,
                user_email=user.email,
                target_type="LEAD",
                target_id=lead.id,
                details=f"Chuyển đổi thành công Lead '{lead.full_name}' -> Khách hàng '{customer.full_name}', Deal ID '{deal.id if deal else 'N/A'}'",
                created_at=now,
            )
            db.add(audit)
            db.commit()

            return {
                "message": "Chuyển đổi đầu mối thành công!",
                "leadId": lead.id,
                "customerId": customer.id,
                "customerName": customer.full_name,
                "contactId": contact.id,
                "dealId": deal.id if deal else None,
            }

        except Exception as exc:
            db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Lỗi khi chuyển đổi đầu mối: {str(exc)}",
            )

    @classmethod
    def check_duplicates_data(
        cls,
        db: Session,
        email: Optional[str] = None,
        phone: Optional[str] = None,
        exclude_lead_id: Optional[str] = None,
    ) -> Dict[str, Any]:
        result = LeadRepository.find_potential_duplicates(
            db, email=email, phone=phone, exclude_lead_id=exclude_lead_id
        )
        leads_dto = [cls.map_to_dto(l) for l in result["leads"]]
        customers_dto = [
            {
                "id": c.id,
                "fullName": c.full_name,
                "email": c.email,
                "phone": c.phone,
                "company": c.company,
                "status": c.status,
            }
            for c in result["customers"]
        ]
        return {
            "duplicateLeads": leads_dto,
            "duplicateCustomers": customers_dto,
            "hasDuplicates": len(leads_dto) > 0 or len(customers_dto) > 0,
        }

    @classmethod
    def merge_leads(cls, db: Session, dto: MergeLeadsDTO, user: User) -> Dict[str, Any]:
        """
        S4-04: Gộp hai hồ sơ Lead trùng lặp với transaction an toàn và ghi lịch sử kiểm toán.
        """
        primary = LeadRepository.get_by_id(db, dto.primaryLeadId)
        duplicate = LeadRepository.get_by_id(db, dto.duplicateLeadId)

        if not primary or not duplicate:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Không tìm thấy một trong hai bản ghi Lead để gộp.",
            )

        if primary.id == duplicate.id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Không thể gộp một bản ghi với chính nó.",
            )

        # Hợp nhất thông tin theo chỉ định của người dùng
        keep_fields = dto.keepFields or {}
        for field, choice in keep_fields.items():
            if choice == "duplicate" and hasattr(duplicate, field):
                setattr(primary, field, getattr(duplicate, field))

        # Hợp nhất ghi chú
        notes_combined = f"{primary.notes or ''}\n[Gộp từ {duplicate.id} lúc {datetime.utcnow().isoformat()}]: {duplicate.notes or ''}".strip()
        primary.notes = notes_combined

        # Cập nhật điểm
        primary.score, primary.score_grade = cls.calculate_score({
            "email": primary.email,
            "phone": primary.phone,
            "company": primary.company,
            "job_title": primary.job_title,
            "company_size": primary.company_size,
            "annual_budget": primary.annual_budget,
        })

        try:
            db.delete(duplicate)
            db.commit()
            db.refresh(primary)

            # Ghi nhật ký kiểm toán
            audit = AuditLog(
                id=str(uuid.uuid4()),
                user_id=user.id,
                action="LEAD_MERGED",
                performed_by=user.full_name,
                user_name=user.full_name,
                user_email=user.email,
                target_type="LEAD",
                target_id=primary.id,
                details=f"Gộp đầu mối '{duplicate.full_name}' ({duplicate.email}) vào hồ sơ chính '{primary.full_name}'.",
                created_at=datetime.utcnow(),
            )
            db.add(audit)
            db.commit()

            return {
                "message": "Gộp bản ghi đầu mối thành công!",
                "primaryLead": cls.map_to_dto(primary),
            }
        except Exception as exc:
            db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Lỗi khi gộp đầu mối: {str(exc)}",
            )

    @classmethod
    def check_sla_overdue(cls, db: Session) -> int:
        """
        S4-07: Quét và cập nhật cờ sla_overdue cho các lead quá hạn phản hồi.
        """
        now = datetime.utcnow()
        overdue_leads = (
            db.query(Lead)
            .filter(
                Lead.status.in_(["UNASSIGNED", "ASSIGNED"]),
                Lead.sla_deadline < now,
                Lead.sla_overdue == False,
            )
            .all()
        )
        count = 0
        for l in overdue_leads:
            l.sla_overdue = True
            count += 1
        if count > 0:
            db.commit()
        return count

    # =========================================================================
    # S4-02: LEAD EXCEL / CSV IMPORT SERVICE
    # =========================================================================

    @staticmethod
    def generate_template(format_type: str = "xlsx") -> StreamingResponse:
        """Sinh file mẫu (.xlsx hoặc .csv) nhập danh sách Leads."""
        if format_type.lower() == "csv":
            output = io.BytesIO()
            output.write(b"\xef\xbb\xbf")  # UTF-8 BOM
            text_stream = io.TextIOWrapper(output, encoding="utf-8", newline="", write_through=True)
            writer = csv.writer(text_stream)
            writer.writerow(["full_name", "email", "phone", "company", "job_title", "region", "industry", "source"])
            writer.writerow(["Nguyễn Văn An", "an.nguyen@congtyabc.vn", "0912345678", "Công ty CP ABC", "Giám Đốc CNTT", "Miền Bắc", "Công nghệ", "Excel Import"])
            writer.writerow(["Trần Thị Bích", "bich.tran@tapdoanxyz.vn", "0987654321", "Tập đoàn XYZ", "Trưởng Phòng Mua Hàng", "Miền Nam", "Tài chính", "Excel Import"])
            text_stream.detach()
            output.seek(0)
            return StreamingResponse(
                output,
                media_type="text/csv",
                headers={"Content-Disposition": 'attachment; filename="leads_import_template.csv"'},
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads_Template"
        headers = ["full_name", "email", "phone", "company", "job_title", "region", "industry", "source"]
        ws.append(headers)
        ws.append(["Nguyễn Văn An", "an.nguyen@congtyabc.vn", "0912345678", "Công ty CP ABC", "Giám Đốc CNTT", "Miền Bắc", "Công nghệ", "Excel Import"])
        ws.append(["Trần Thị Bích", "bich.tran@tapdoanxyz.vn", "0987654321", "Tập đoàn XYZ", "Trưởng Phòng Mua Hàng", "Miền Nam", "Tài chính", "Excel Import"])

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="leads_import_template.xlsx"'},
        )

    @classmethod
    def execute_bulk_import(cls, db: Session, raw_rows: List[Dict[str, Any]], user: User) -> Dict[str, Any]:
        """
        S4-02: Nhập hàng loạt Leads từ dữ liệu đã parse kèm kiểm tra trùng lặp và tính điểm tự động.
        """
        success_count = 0
        error_count = 0
        errors = []

        existing_emails = {l[0].lower() for l in db.query(Lead.email).all()}

        for idx, row in enumerate(raw_rows, start=2):
            name = (row.get("full_name") or row.get("fullname") or row.get("name") or "").strip()
            email = (row.get("email") or "").strip().lower()
            phone = (row.get("phone") or "").strip()
            company = (row.get("company") or "").strip()
            job_title = (row.get("job_title") or row.get("title") or "").strip()
            region = (row.get("region") or "Miền Bắc").strip()
            industry = (row.get("industry") or "Công nghệ").strip()
            source = (row.get("source") or "Excel Import").strip()

            if not name or not email:
                error_count += 1
                errors.append({"row": idx, "email": email, "reason": "Thiếu Họ tên hoặc Email bắt buộc."})
                continue

            if "@" not in email:
                error_count += 1
                errors.append({"row": idx, "email": email, "reason": "Định dạng Email không hợp lệ."})
                continue

            if email in existing_emails:
                error_count += 1
                errors.append({"row": idx, "email": email, "reason": "Email đã tồn tại trong danh sách Đầu mối."})
                continue

            score, grade = cls.calculate_score({
                "email": email,
                "phone": phone,
                "company": company,
                "job_title": job_title,
            })

            now = datetime.utcnow()
            lead = Lead(
                id=str(uuid.uuid4()),
                full_name=name,
                email=email,
                phone=phone or None,
                company=company or None,
                job_title=job_title or None,
                source=source,
                region=region,
                industry=industry,
                score=score,
                score_grade=grade,
                status="UNASSIGNED",
                sla_deadline=now + timedelta(hours=24),
                sla_overdue=False,
                created_at=now,
                updated_at=now,
            )

            # Auto-route
            auto_u = cls.auto_route_lead(db, lead)
            if auto_u:
                lead.assigned_user_id = auto_u.id
                lead.assigned_at = now
                lead.status = "ASSIGNED"

            db.add(lead)
            existing_emails.add(email)
            success_count += 1

        db.commit()

        return {
            "totalRows": len(raw_rows),
            "successCount": success_count,
            "errorCount": error_count,
            "total_rows": len(raw_rows),
            "success_count": success_count,
            "error_count": error_count,
            "errors": errors,
        }
