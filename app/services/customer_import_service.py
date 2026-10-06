"""
Tầng Service xử lý nghiệp vụ Nhập Danh Sách Khách Hàng Hàng Loạt từ Excel/CSV (SCRUM-64 / S3-06).
Bao gồm:
- Tạo và tải tệp Excel mẫu chuẩn (.xlsx) kèm dữ liệu mẫu doanh nghiệp.
- Xem trước (Preview) kiểm tra hợp lệ từng dòng, bắt lỗi chi tiết và quét trùng lặp CSDL (tax_code, website).
- Thực thi nhập dữ liệu hàng loạt (Execute Import) với lựa chọn xử lý trùng lặp (SKIP / UPDATE) và transaction an toàn.
"""
import io
import csv
import re
import uuid
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional, Set
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import or_
from fastapi import HTTPException, status, UploadFile
from fastapi.responses import StreamingResponse

from app.models.customer import Customer
from app.models.user import User
from app.core.customer_search import normalize_phone, normalize_tax_code, normalize_text
from app.schemas.customer_import import (
    CustomerImportRow,
    CustomerImportPreviewResponse,
    CustomerImportExecuteRequest,
    CustomerImportSummaryResponse,
)

# Regex chuẩn kiểm tra dữ liệu
EMAIL_REGEX = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
VN_PHONE_REGEX = re.compile(r"^(?:0|\+84)(3|5|7|8|9)\d{8}$")
TAX_CODE_REGEX = re.compile(r"^\d{10}(?:-?\d{3})?$")

# Ánh xạ trạng thái khách hàng sang CSDL
STATUS_MAPPING: Dict[str, str] = {
    "LEAD": "lead",
    "CONTACTED": "prospect",
    "CUSTOMER": "active",
    "ACTIVE": "active",
    "INACTIVE": "inactive",
    "PROSPECT": "prospect",
}

# Danh mục tiêu đề cột linh hoạt hỗ trợ cả tiếng Anh và tiếng Việt
HEADER_SYNONYMS: Dict[str, str] = {
    # Tên khách hàng / doanh nghiệp
    "name": "name",
    "tên": "name",
    "ten": "name",
    "tên khách hàng": "name",
    "ten khach hang": "name",
    "họ và tên": "name",
    "ho va ten": "name",
    "tên công ty": "name",
    "ten cong ty": "name",
    "công ty": "name",
    "cong ty": "name",
    "doanh nghiệp": "name",
    "doanh nghiep": "name",
    "customer name": "name",
    "company": "name",
    # Mã số thuế
    "tax_code": "tax_code",
    "tax code": "tax_code",
    "taxcode": "tax_code",
    "mã số thuế": "tax_code",
    "ma so thue": "tax_code",
    "mst": "tax_code",
    # Email
    "email": "email",
    "hòm thư": "email",
    "hom thu": "email",
    "mail": "email",
    # Số điện thoại
    "phone": "phone",
    "số điện thoại": "phone",
    "so dien thoai": "phone",
    "điện thoại": "phone",
    "dien thoai": "phone",
    "sđt": "phone",
    "sdt": "phone",
    "hotline": "phone",
    # Website
    "website": "website",
    "web": "website",
    "domain": "website",
    "tên miền": "website",
    "ten mien": "website",
    "trang web": "website",
    "url": "website",
    # Địa chỉ
    "address": "address",
    "địa chỉ": "address",
    "dia chi": "address",
    "địa điểm": "address",
    "dia diem": "address",
    "khu vực": "address",
    "khu vuc": "address",
    "location": "address",
    # Trạng thái
    "status": "status",
    "trạng thái": "status",
    "trang thai": "status",
}


def clean_website_url(url: Optional[str]) -> str:
    """Chuẩn hóa website để so sánh đối soát trùng lặp."""
    if not url:
        return ""
    w = url.strip().lower()
    w = re.sub(r"^https?://", "", w)
    w = re.sub(r"^www\.", "", w)
    return w.rstrip("/")


class CustomerImportService:
    """
    Service cung cấp logic nghiệp vụ cho SCRUM-64:
    - Tải template
    - Xem trước (preview) & validate
    - Thực thi nạp CSDL (execute)
    """

    @staticmethod
    def generate_template() -> StreamingResponse:
        """
        Sinh tệp Excel mẫu chuẩn (.xlsx) gồm 7 cột:
        name, tax_code, email, phone, website, address, status
        Kèm định dạng đẹp và 2 dòng dữ liệu mẫu doanh nghiệp.
        """
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Customer_Import_Template"

        headers = ["name", "tax_code", "email", "phone", "website", "address", "status"]
        sample_rows = [
            [
                "Tập đoàn Công nghệ FPT",
                "0101248141",
                "contact@fpt.com.vn",
                "02473007300",
                "https://fpt.com.vn",
                "Số 10 Phạm Văn Bạch, Cầu Giấy, Hà Nội",
                "CUSTOMER",
            ],
            [
                "Công ty TNHH Vận tải Toàn Cầu",
                "0312345678",
                "info@toancau.vn",
                "0901234567",
                "https://toancau.vn",
                "123 Nguyễn Văn Linh, Quận 7, TP.HCM",
                "LEAD",
            ],
        ]

        # Định dạng Header
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        worksheet.row_dimensions[1].height = 28
        worksheet.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Dữ liệu mẫu
        data_font = Font(name="Arial", size=10)
        data_align = Alignment(vertical="center")
        for row_data in sample_rows:
            worksheet.append(row_data)

        for row_idx in range(2, len(sample_rows) + 2):
            worksheet.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(headers) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border

        # Thiết lập độ rộng cột tự động
        col_widths = {
            1: 32,  # name
            2: 18,  # tax_code
            3: 26,  # email
            4: 16,  # phone
            5: 25,  # website
            6: 38,  # address
            7: 15,  # status
        }
        for col_idx, width in col_widths.items():
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = width

        output_stream = io.BytesIO()
        workbook.save(output_stream)
        output_stream.seek(0)

        filename = "customer_import_template.xlsx"
        return StreamingResponse(
            output_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    @classmethod
    def _parse_uploaded_file(cls, content: bytes, filename: str) -> List[Dict[str, Any]]:
        """
        Đọc nội dung tệp .xlsx hoặc .csv và chuyển thành danh sách từ điển theo cột chuẩn hóa.
        """
        filename_lower = (filename or "").lower()
        rows_data: List[Dict[str, Any]] = []

        if filename_lower.endswith(".xlsx") or filename_lower.endswith(".xls"):
            try:
                wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
                ws = wb.active
                if ws is None:
                    return []
                all_rows = list(ws.iter_rows(values_only=True))
                if not all_rows:
                    return []

                raw_headers = [str(cell or "").strip() for cell in all_rows[0]]
                mapped_headers = [HEADER_SYNONYMS.get(h.lower(), h.lower()) for h in raw_headers]

                for row_idx, row_values in enumerate(all_rows[1:], start=2):
                    # Bỏ qua dòng hoàn toàn rỗng
                    if not any(v is not None and str(v).strip() != "" for v in row_values):
                        continue
                    row_dict: Dict[str, Any] = {"_row_number": row_idx}
                    for h_idx, col_key in enumerate(mapped_headers):
                        if h_idx < len(row_values):
                            val = row_values[h_idx]
                            row_dict[col_key] = str(val).strip() if val is not None else ""
                    rows_data.append(row_dict)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Không thể đọc tệp Excel. Tệp có thể bị hỏng hoặc sai định dạng: {str(e)}",
                )

        elif filename_lower.endswith(".csv"):
            text_content = ""
            for encoding in ["utf-8-sig", "utf-8", "latin-1"]:
                try:
                    text_content = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if not text_content:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Không thể giải mã tệp CSV. Vui lòng lưu tệp ở định dạng UTF-8.",
                )

            reader = csv.reader(io.StringIO(text_content))
            all_rows = list(reader)
            if not all_rows:
                return []

            raw_headers = [h.strip() for h in all_rows[0]]
            mapped_headers = [HEADER_SYNONYMS.get(h.lower(), h.lower()) for h in raw_headers]

            for row_idx, row_values in enumerate(all_rows[1:], start=2):
                if not any(v.strip() != "" for v in row_values):
                    continue
                row_dict: Dict[str, Any] = {"_row_number": row_idx}
                for h_idx, col_key in enumerate(mapped_headers):
                    if h_idx < len(row_values):
                        row_dict[col_key] = row_values[h_idx].strip()
                rows_data.append(row_dict)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Chỉ chấp nhận tệp định dạng Excel (.xlsx) hoặc CSV (.csv).",
            )

        return rows_data

    @classmethod
    async def preview_import(
        cls,
        file: UploadFile,
        db: Session,
        current_user: User,
    ) -> CustomerImportPreviewResponse:
        """
        Đọc và kiểm tra tính hợp lệ dữ liệu (dry-run preview):
        - Bắt lỗi dữ liệu từng dòng
        - Phát hiện trùng lặp CSDL (tax_code, website) và trùng lặp nội bộ file
        - Trả về danh sách preview kèm trạng thái
        """
        # Kiểm tra kích thước tệp tối đa 5MB
        content = await file.read()
        if len(content) > 5 * 1024 * 1024:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="Kích thước tệp vượt quá giới hạn cho phép (tối đa 5MB).",
            )
        if len(content) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Tệp tải lên rỗng không có nội dung.",
            )

        raw_rows = cls._parse_uploaded_file(content, file.filename or "upload.xlsx")
        if not raw_rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Tệp tải lên không chứa dòng dữ liệu nào hoặc thiếu tiêu đề cột hợp lệ.",
            )

        # 1. Gom danh sách tax_code và website để Batch Query kiểm tra trùng CSDL
        tax_codes_in_file: Set[str] = set()
        websites_in_file: Set[str] = set()

        for r in raw_rows:
            tc = (r.get("tax_code") or "").strip()
            if tc:
                norm_tc = normalize_tax_code(tc)
                if norm_tc:
                    tax_codes_in_file.add(norm_tc)
                tax_codes_in_file.add(tc)

            ws = clean_website_url(r.get("website"))
            if ws:
                websites_in_file.add(ws)

        # 2. Truy vấn khách hàng hiện có trong CSDL (chưa bị xóa mềm)
        existing_customers: List[Customer] = []
        if tax_codes_in_file or websites_in_file:
            query = db.query(Customer).filter(Customer.is_deleted == False)
            conditions = []
            if tax_codes_in_file:
                conditions.append(Customer.tax_code.in_(list(tax_codes_in_file)))
                conditions.append(Customer.normalized_tax_code.in_(list(tax_codes_in_file)))
            if websites_in_file:
                conditions.append(Customer.website.in_(list(websites_in_file)))
            if conditions:
                query = query.filter(or_(*conditions))
            existing_customers = query.all()

        # Tạo map tra cứu nhanh O(1)
        db_by_tax_code: Dict[str, Customer] = {}
        db_by_website: Dict[str, Customer] = {}

        for c in existing_customers:
            if c.tax_code:
                db_by_tax_code[c.tax_code.strip()] = c
            if c.normalized_tax_code:
                db_by_tax_code[c.normalized_tax_code.strip()] = c
            if c.website:
                clean_w = clean_website_url(c.website)
                if clean_w:
                    db_by_website[clean_w] = c

        # 3. Duyệt từng dòng để Validate & Check Duplicate
        processed_rows: List[CustomerImportRow] = []
        seen_tax_in_file: Dict[str, int] = {}
        seen_website_in_file: Dict[str, int] = {}

        for r in raw_rows:
            row_num = r.get("_row_number", len(processed_rows) + 2)
            name = (r.get("name") or "").strip()
            tax_code = (r.get("tax_code") or "").strip() or None
            email = (r.get("email") or "").strip() or None
            phone = (r.get("phone") or "").strip() or None
            website = (r.get("website") or "").strip() or None
            address = (r.get("address") or "").strip() or None
            raw_status = (r.get("status") or "LEAD").strip().upper()

            errors: List[str] = []

            # Validate name: bắt buộc, 2 - 150 ký tự
            if not name:
                errors.append("Tên khách hàng là bắt buộc.")
            elif len(name) < 2 or len(name) > 150:
                errors.append("Tên khách hàng phải từ 2 đến 150 ký tự.")

            # Validate tax_code: nếu có thì phải 10 hoặc 13 chữ số
            if tax_code:
                clean_tc = tax_code.replace("-", "").strip()
                if not (clean_tc.isdigit() and len(clean_tc) in (10, 13)):
                    errors.append("Mã số thuế không đúng định dạng (phải gồm 10 hoặc 13 chữ số).")

            # Validate email: nếu có phải đúng regex
            if email:
                if not EMAIL_REGEX.match(email):
                    errors.append(f"Email không đúng định dạng: '{email}'.")

            # Validate phone: nếu có phải đúng định dạng số Việt Nam 10 chữ số
            if phone:
                clean_p = phone.replace(" ", "").replace(".", "").replace("-", "")
                if not (VN_PHONE_REGEX.match(clean_p) or (clean_p.isdigit() and len(clean_p) == 10)):
                    errors.append(f"Số điện thoại không đúng định dạng Việt Nam: '{phone}'.")

            # Validate status
            if raw_status not in STATUS_MAPPING:
                raw_status = "LEAD"

            # Kiểm tra trùng lặp
            is_duplicate = False
            duplicate_reason: Optional[str] = None
            matched_id: Optional[str] = None
            matched_name: Optional[str] = None

            norm_tc = normalize_tax_code(tax_code or "")
            clean_ws = clean_website_url(website)

            # A. Kiểm tra với CSDL
            if tax_code and (tax_code in db_by_tax_code or (norm_tc and norm_tc in db_by_tax_code)):
                matched = db_by_tax_code.get(tax_code) or db_by_tax_code.get(norm_tc)
                if matched:
                    is_duplicate = True
                    matched_id = matched.id
                    matched_name = matched.full_name or matched.company
                    duplicate_reason = f"Trùng mã số thuế '{tax_code}' với khách hàng '{matched_name}' trong hệ thống"

            if not is_duplicate and clean_ws and clean_ws in db_by_website:
                matched = db_by_website[clean_ws]
                is_duplicate = True
                matched_id = matched.id
                matched_name = matched.full_name or matched.company
                duplicate_reason = f"Trùng địa chỉ website '{website}' với khách hàng '{matched_name}' trong hệ thống"

            # B. Kiểm tra trùng lặp nội bộ trong file tải lên
            if not is_duplicate:
                if norm_tc and norm_tc in seen_tax_in_file:
                    is_duplicate = True
                    prev_row = seen_tax_in_file[norm_tc]
                    duplicate_reason = f"Trùng mã số thuế với dòng {prev_row} trong cùng tệp tải lên"
                elif clean_ws and clean_ws in seen_website_in_file:
                    is_duplicate = True
                    prev_row = seen_website_in_file[clean_ws]
                    duplicate_reason = f"Trùng địa chỉ website với dòng {prev_row} trong cùng tệp tải lên"

            if norm_tc:
                seen_tax_in_file[norm_tc] = row_num
            if clean_ws:
                seen_website_in_file[clean_ws] = row_num

            is_valid = len(errors) == 0

            row_dto = CustomerImportRow(
                row_number=row_num,
                name=name,
                tax_code=tax_code,
                email=email,
                phone=phone,
                website=website,
                address=address,
                status=raw_status,
                is_valid=is_valid,
                errors=errors,
                is_duplicate=is_duplicate,
                duplicate_reason=duplicate_reason,
                matched_customer_id=matched_id,
                matched_customer_name=matched_name,
            )
            processed_rows.append(row_dto)

        valid_count = sum(1 for r in processed_rows if r.is_valid)
        invalid_count = sum(1 for r in processed_rows if not r.is_valid)
        duplicate_count = sum(1 for r in processed_rows if r.is_duplicate)

        return CustomerImportPreviewResponse(
            total_rows=len(processed_rows),
            valid_rows_count=valid_count,
            invalid_rows_count=invalid_count,
            duplicate_rows_count=duplicate_count,
            rows=processed_rows,
        )

    @classmethod
    def execute_import(
        cls,
        request: CustomerImportExecuteRequest,
        db: Session,
        current_user: User,
    ) -> CustomerImportSummaryResponse:
        """
        Thực thi nhập dữ liệu hàng loạt vào CSDL:
        - Bỏ qua các dòng có is_valid == False
        - Nếu is_duplicate == True:
          + "SKIP": Bỏ qua, tăng skipped_count
          + "UPDATE": Cập nhật thông tin mới vào khách hàng đã có, tăng updated_count
        - Nếu is_duplicate == False: Tạo khách hàng mới, tăng imported_count
        - Đảm bảo Transaction an toàn (commit / rollback)
        """
        total_rows = len(request.rows)
        imported_count = 0
        updated_count = 0
        skipped_count = 0
        failed_count = 0
        errors_detail: List[Dict[str, Any]] = []

        try:
            for row in request.rows:
                # 1. Cơ chế bỏ qua lỗi: dòng có lỗi validation tự động bị loại
                if not row.is_valid:
                    failed_count += 1
                    errors_detail.append({
                        "row_number": row.row_number,
                        "name": row.name,
                        "errors": row.errors,
                        "reason": "; ".join(row.errors) if row.errors else "Dữ liệu dòng không hợp lệ.",
                    })
                    continue

                # 2. Xử lý bản ghi trùng lặp
                if row.is_duplicate:
                    if request.duplicate_action.upper() == "SKIP":
                        skipped_count += 1
                        continue

                    if request.duplicate_action.upper() == "UPDATE":
                        # Tìm bản ghi khách hàng hiện có
                        target_customer: Optional[Customer] = None
                        if row.matched_customer_id:
                            target_customer = db.query(Customer).filter(
                                Customer.id == row.matched_customer_id,
                                Customer.is_deleted == False
                            ).first()

                        if not target_customer and row.tax_code:
                            target_customer = db.query(Customer).filter(
                                Customer.tax_code == row.tax_code,
                                Customer.is_deleted == False
                            ).first()

                        if target_customer:
                            # Cập nhật thông tin mới
                            if row.name:
                                target_customer.full_name = row.name
                                target_customer.company = row.name
                            if row.email:
                                target_customer.email = row.email
                            if row.phone:
                                target_customer.phone = row.phone
                            if row.tax_code:
                                target_customer.tax_code = row.tax_code
                            if row.website:
                                target_customer.website = row.website
                            if row.address:
                                target_customer.address = row.address
                                target_customer.region = row.address
                            if row.status:
                                target_customer.status = STATUS_MAPPING.get(row.status.upper(), "lead")
                            target_customer.updated_at = datetime.utcnow()
                            updated_count += 1
                            continue
                        else:
                            # Nếu không tìm thấy bản ghi cũ, tạo mới an toàn
                            pass

                # 3. Tạo mới khách hàng
                assigned_id = current_user.id
                db_status = STATUS_MAPPING.get((row.status or "LEAD").upper(), "lead")

                # Fallback email và phone hợp lệ nếu chưa có
                gen_email = row.email or f"customer_{uuid.uuid4().hex[:8]}@unassigned.vn"
                gen_phone = row.phone or "0900000000"

                new_customer = Customer(
                    id=str(uuid.uuid4()),
                    full_name=row.name,
                    company=row.name,
                    email=gen_email,
                    phone=gen_phone,
                    tax_code=row.tax_code,
                    website=row.website,
                    address=row.address,
                    region=row.address,
                    status=db_status,
                    health_score=85,
                    assigned_user_id=assigned_id,
                    is_deleted=False,
                    created_at=datetime.utcnow(),
                    updated_at=datetime.utcnow(),
                )
                db.add(new_customer)
                imported_count += 1

            # Commit toàn bộ trong một Transaction an toàn
            db.commit()

        except Exception as e:
            db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Đã xảy ra lỗi trong quá trình thực thi nhập dữ liệu: {str(e)}",
            )

        return CustomerImportSummaryResponse(
            total_rows=total_rows,
            imported_count=imported_count,
            updated_count=updated_count,
            skipped_count=skipped_count,
            failed_count=failed_count,
            errors_detail=errors_detail,
        )
