import io
import json
import csv
from datetime import datetime
from typing import List, Optional, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, status, Query, Response, UploadFile, File, Request
from sqlalchemy.orm import Session
import openpyxl

from app.database import get_db
from app.dependencies import get_current_user
from app.models.user import User
from app.models.lead import Lead
from app.schemas.lead import (
    LeadDTO,
    CreateLeadDTO,
    UpdateLeadDTO,
    PublicLeadSubmissionDTO,
    RejectLeadDTO,
    MergeLeadsDTO,
    ConvertLeadDTO,
)
from app.services.lead_service import LeadService

router = APIRouter(prefix="/leads", tags=["Lead Management (EP-04: S4-01 -> S4-09)"])


# ==============================================================================
# S4-01: PUBLIC LEAD CAPTURE & EMBED CONFIGURATION
# ==============================================================================

@router.post(
    "/public/submit",
    response_model=LeadDTO,
    status_code=status.HTTP_201_CREATED,
    summary="Tiếp nhận đầu mối công khai từ Website / Form nhúng (S4-01)",
)
def submit_public_lead(
    dto: PublicLeadSubmissionDTO,
    request: Request,
    db: Session = Depends(get_db),
) -> LeadDTO:
    """
    S4-01: Endpoint công khai cho phép nhúng form thu thập thông tin khách hàng tiềm năng:
    - Không cần Header Authorization.
    - Tự động gắn nguồn 'Web Form'.
    - Tự động chấm điểm ban đầu và phân bổ Round-Robin/Khu vực.
    - Khởi tạo SLA phản hồi trong 24 giờ.
    """
    lead = LeadService.submit_public_lead(db, dto)
    return LeadService.map_to_dto(lead)


@router.get(
    "/embed-snippet",
    summary="Lấy mã nhúng HTML iFrame / JavaScript Form thu thập Lead (S4-01)",
)
@router.get(
    "/public/embed-snippet",
    summary="Lấy mã nhúng HTML iFrame / JavaScript Form thu thập Lead (S4-01)",
)
def get_embed_snippet(
    campaign_code: Optional[str] = Query(None, description="Mã chiến dịch tự động gắn khi khách gửi form"),
):
    """
    S4-01: Cung cấp đoạn mã nhúng iFrame hoặc JavaScript Widget cho bộ phận Marketing
    sao chép vào website hoặc landing page mà không để lộ secret API key.
    """
    campaign_param = f"?campaign={campaign_code}" if campaign_code else ""
    iframe_code = (
        f'<iframe src="https://crm.nexus.vn/embed/lead-form{campaign_param}" '
        f'width="100%" height="480" frameborder="0" style="border:none; border-radius:12px; '
        f'box-shadow:0 4px 20px rgba(0,0,0,0.08);" title="NexusCRM Lead Form"></iframe>'
    )
    js_code = (
        f'<div id="nexus-lead-form"></div>\n'
        f'<script src="https://crm.nexus.vn/static/lead-widget.js" '
        f'data-campaign="{campaign_code or "DEFAULT"}" async></script>'
    )
    return {
        "embed_html": iframe_code,
        "iframeSnippet": iframe_code,
        "script_tag": js_code,
        "scriptSnippet": js_code,
        "form_action_url": "/api/v1/leads/public/submit",
        "targetEndpoint": "/api/v1/leads/public/submit",
        "campaign_code": campaign_code or "DEFAULT",
        "description": "Sao chép đoạn mã trên và dán vào trang web hoặc Landing Page của doanh nghiệp.",
    }


# ==============================================================================
# S4-02: EXCEL / CSV IMPORT CHO LEADS
# ==============================================================================

@router.get("/import/template", summary="Tải tệp mẫu Excel hoặc CSV nhập danh sách Đầu mối (S4-02)")
def download_lead_template(
    format: str = Query("xlsx", description="xlsx hoặc csv"),
    current_user: User = Depends(get_current_user),
):
    return LeadService.generate_template(format_type=format)


@router.post("/import/preview", summary="Xem trước tệp Excel/CSV nhập danh sách Đầu mối (S4-02)")
async def preview_lead_import(
    file: UploadFile = File(..., description="Tệp Excel (.xlsx, .xls) hoặc CSV (.csv)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    file_bytes = await file.read()
    filename = (file.filename or "").lower()

    rows = []
    if filename.endswith(".csv"):
        text_content = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text_content))
        for r in reader:
            rows.append(r)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if all_rows:
            headers = [str(h or "").strip().lower() for h in all_rows[0]]
            for r_data in all_rows[1:]:
                row_dict = {}
                for idx, val in enumerate(r_data):
                    if idx < len(headers):
                        row_dict[headers[idx]] = str(val or "").strip()
                if any(row_dict.values()):
                    rows.append(row_dict)

    existing_emails = {l[0].lower() for l in db.query(Lead.email).all()}
    preview_details = []
    valid_count = 0
    error_count = 0

    for idx, r in enumerate(rows, start=2):
        name = (r.get("full_name") or r.get("fullname") or r.get("name") or "").strip()
        email = (r.get("email") or "").strip().lower()
        phone = (r.get("phone") or "").strip()
        company = (r.get("company") or "").strip()
        errors = []

        if not name:
            errors.append("Thiếu Họ và tên.")
        if not email:
            errors.append("Thiếu Email.")
        elif "@" not in email:
            errors.append("Định dạng Email không hợp lệ.")
        elif email in existing_emails:
            errors.append("Email đã tồn tại trong danh sách Đầu mối.")

        if errors:
            error_count += 1
            status_val = "INVALID"
        else:
            valid_count += 1
            status_val = "VALID"

        preview_details.append({
            "rowIndex": idx,
            "fullName": name,
            "email": email,
            "phone": phone,
            "company": company,
            "status": status_val,
            "errors": errors,
        })

    return {
        "totalRows": len(rows),
        "validCount": valid_count,
        "errorCount": error_count,
        "details": preview_details,
    }


@router.post("/import-excel", summary="Nhập hàng loạt Leads từ file Excel/CSV (S4-02)")
@router.post("/import/execute", summary="Thực thi nhập hàng loạt Đầu mối từ file Excel/CSV hoặc JSON (S4-02)")
async def execute_lead_import(
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    content_type = request.headers.get("content-type", "").lower()
    raw_rows = []

    if "multipart/form-data" in content_type:
        form = await request.form()
        file_obj = form.get("file")
        if not file_obj or not hasattr(file_obj, "read"):
            raise HTTPException(status_code=400, detail="Vui lòng đính kèm tệp 'file'.")
        file_bytes = await file_obj.read()
        filename = (getattr(file_obj, "filename", "") or "").lower()

        if filename.endswith(".csv"):
            text_content = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text_content))
            for r in reader:
                raw_rows.append(r)
        else:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if all_rows:
                headers = [str(h or "").strip().lower() for h in all_rows[0]]
                for r_data in all_rows[1:]:
                    row_dict = {}
                    for idx, val in enumerate(r_data):
                        if idx < len(headers):
                            row_dict[headers[idx]] = str(val or "").strip()
                    if any(row_dict.values()):
                        raw_rows.append(row_dict)
    else:
        try:
            body = await request.json()
            raw_rows = body.get("rows") or body if isinstance(body, list) else []
        except Exception:
            raise HTTPException(status_code=400, detail="Dữ liệu JSON không hợp lệ.")

    if not raw_rows:
        raise HTTPException(status_code=400, detail="Không có dòng dữ liệu nào để nhập.")

    return LeadService.execute_bulk_import(db=db, raw_rows=raw_rows, user=current_user)


# ==============================================================================
# S4-06: ROUTING RULES
# ==============================================================================

@router.get("/routing-rules", summary="Lấy danh sách quy tắc phân bổ định tuyến Lead (S4-06)")
def get_routing_rules(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from app.repositories.lead_repository import LeadRepository
    rules = LeadRepository.get_routing_rules(db)
    results = []
    for r in rules:
        target_name = None
        if r.target_user_id:
            u = db.query(User).filter(User.id == r.target_user_id).first()
            if u:
                target_name = u.full_name
        results.append({
            "id": r.id,
            "name": r.name,
            "rule_type": r.rule_type,
            "ruleType": r.rule_type,
            "criteria_value": r.criterion_value,
            "criterionValue": r.criterion_value,
            "assigned_user_id": r.target_user_id,
            "targetUserId": r.target_user_id,
            "assigned_user_name": target_name,
            "targetUserName": target_name,
            "priority": r.priority,
            "is_active": r.is_active,
            "isActive": r.is_active,
            "created_at": r.created_at.isoformat() if r.created_at else None,
        })
    return results


@router.post("/routing-rules", summary="Tạo mới quy tắc phân bổ định tuyến Lead (S4-06)")
def create_routing_rule(
    dto: Dict[str, Any],
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    import uuid
    from app.models.lead import LeadRoutingRule
    new_rule = LeadRoutingRule(
        id=str(uuid.uuid4()),
        name=dto.get("name") or "Quy tắc định tuyến mới",
        priority=int(dto.get("priority") or 1),
        rule_type=dto.get("ruleType") or dto.get("rule_type") or "REGION",
        criterion_value=dto.get("criterionValue") or dto.get("criteria_value"),
        target_user_id=dto.get("targetUserId") or dto.get("assigned_user_id"),
        is_active=bool(dto.get("isActive", dto.get("is_active", True))),
    )
    db.add(new_rule)
    db.commit()
    db.refresh(new_rule)
    return {
        "id": new_rule.id,
        "name": new_rule.name,
        "rule_type": new_rule.rule_type,
        "criteria_value": new_rule.criterion_value,
        "assigned_user_id": new_rule.target_user_id,
        "priority": new_rule.priority,
        "is_active": new_rule.is_active,
    }


# ==============================================================================
# S4-04: DUPLICATE DETECTION & MERGE WORKFLOW
# ==============================================================================

@router.get("/duplicates", summary="Phát hiện trùng lặp đầu mối theo Email hoặc SĐT (S4-04)")
def check_duplicates(
    email: Optional[str] = Query(None, description="Địa chỉ Email cần kiểm tra"),
    phone: Optional[str] = Query(None, description="Số điện thoại cần kiểm tra"),
    exclude_id: Optional[str] = Query(None, description="Bỏ qua ID lead hiện tại"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    S4-04: So sánh đối soát với cả bảng leads và bảng customers để tìm bản ghi trùng lặp.
    """
    dups = LeadService.check_duplicates_data(db, email=email, phone=phone, exclude_lead_id=exclude_id)
    return dups


@router.get("/{lead_id}/duplicates", summary="Kiểm tra trùng lặp cho một Lead cụ thể (S4-04)")
def check_lead_duplicates_by_id(
    lead_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead.")
    dups_data = LeadService.check_duplicates_data(db, email=lead.email, phone=lead.phone, exclude_lead_id=lead_id)

    # Format list according to frontend ILeadDuplicateMatch interface
    match_list = []
    for l in dups_data.get("duplicateLeads", []):
        matched_field = "Email" if l.email and lead.email and l.email.lower() == lead.email.lower() else "Số điện thoại"
        match_list.append({
            "match_type": "LEAD",
            "record_id": l.id,
            "matched_field": f"{matched_field} ({l.email if matched_field == 'Email' else l.phone})",
            "record_name": l.fullName or l.full_name or "Đầu mối trùng",
            "record_email": l.email,
            "record_phone": l.phone,
            "created_at": l.createdAt or datetime.utcnow().isoformat(),
        })
    for c in dups_data.get("duplicateCustomers", []):
        matched_field = "Email" if c.get("email") and lead.email and c.get("email").lower() == lead.email.lower() else "Số điện thoại"
        match_list.append({
            "match_type": "CUSTOMER",
            "record_id": c.get("id"),
            "matched_field": f"{matched_field} ({c.get('email') if matched_field == 'Email' else c.get('phone')})",
            "record_name": c.get("fullName") or "Khách hàng trùng",
            "record_email": c.get("email"),
            "record_phone": c.get("phone"),
            "created_at": datetime.utcnow().isoformat(),
        })

    return {
        "duplicates": match_list,
        "count": len(match_list),
        "raw": dups_data,
    }


@router.post("/merge", summary="Gộp hai hồ sơ Đầu mối trùng lặp (S4-04)")
def merge_leads_direct(
    dto: MergeLeadsDTO,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return LeadService.merge_leads(db=db, dto=dto, user=current_user)


@router.post("/{lead_id}/merge", summary="Gộp hai hồ sơ Đầu mối trùng lặp (S4-04)")
def merge_leads(
    lead_id: str,
    dto: MergeLeadsDTO,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    S4-04: Thực hiện gộp hồ sơ trùng lặp với transaction an toàn và ghi lịch sử kiểm toán.
    """
    if dto.primaryLeadId != lead_id:
        dto.primaryLeadId = lead_id
    return LeadService.merge_leads(db=db, dto=dto, user=current_user)


# ==============================================================================
# S4-05: LEAD SCORING RECALCULATION
# ==============================================================================

@router.post("/{lead_id}/recalculate-score", summary="Tính toán lại điểm số Lead alias (S4-05)")
@router.post("/{lead_id}/score/recalculate", response_model=LeadDTO, summary="Tính toán lại điểm số Lead (S4-05)")
def recalculate_score(
    lead_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    """
    S4-05: Chấm lại điểm số và phân loại Hot/Warm/Cold dựa trên thuộc tính mới nhất.
    """
    lead_model = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead_model:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead.")

    score, grade = LeadService.calculate_score({
        "email": lead_model.email,
        "phone": lead_model.phone,
        "company": lead_model.company,
        "job_title": lead_model.job_title,
        "company_size": lead_model.company_size,
        "annual_budget": lead_model.annual_budget,
    })
    lead_model.score = score
    lead_model.score_grade = grade
    db.commit()
    db.refresh(lead_model)
    return LeadService.map_to_dto(lead_model)


# ==============================================================================
# S4-06: AUTOMATIC ROUTING
# ==============================================================================

@router.post("/{lead_id}/auto-route", response_model=LeadDTO, summary="Kích hoạt tự động phân bổ Lead (S4-06)")
def auto_route_lead(
    lead_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    lead_model = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead_model:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead.")

    target_user = LeadService.auto_route_lead(db, lead_model)
    if target_user:
        lead_model.assigned_user_id = target_user.id
        lead_model.assigned_at = datetime.utcnow()
        if lead_model.status == "UNASSIGNED":
            lead_model.status = "ASSIGNED"
        db.commit()
        db.refresh(lead_model)

    return LeadService.map_to_dto(lead_model)


# ==============================================================================
# S4-07: ACCEPTANCE, REJECTION & SLA WORKFLOW
# ==============================================================================

@router.post("/{lead_id}/accept", response_model=LeadDTO, summary="Tiếp nhận Lead (S4-07)")
def accept_lead(
    lead_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    return LeadService.accept_lead(db=db, lead_id=lead_id, user=current_user)


@router.post("/{lead_id}/reject", response_model=LeadDTO, summary="Từ chối Lead kèm lý do bắt buộc (S4-07)")
def reject_lead(
    lead_id: str,
    dto: RejectLeadDTO,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    return LeadService.reject_lead(db=db, lead_id=lead_id, dto=dto, user=current_user)


@router.post("/check-sla", summary="Kiểm tra và cập nhật trạng thái vi phạm SLA (S4-07)")
def check_sla(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    overdue_count = LeadService.check_sla_overdue(db)
    return {"message": f"Đã quét và cập nhật {overdue_count} đầu mối quá hạn SLA.", "overdueCount": overdue_count}


# ==============================================================================
# S4-08: LEAD CONVERSION WORKFLOW
# ==============================================================================

@router.post("/{lead_id}/convert", summary="Chuyển đổi Lead thành Customer, Contact và Deal (S4-08)")
def convert_lead(
    lead_id: str,
    dto: ConvertLeadDTO,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return LeadService.convert_lead(db=db, lead_id=lead_id, dto=dto, user=current_user)


# ==============================================================================
# S4-09: LEAD LIST & CRUD ENDPOINTS
# ==============================================================================

@router.get("", response_model=List[LeadDTO], summary="Lấy danh sách Đầu mối bán hàng có phân trang và bộ lọc (S4-09)")
def get_leads(
    response: Response,
    search: Optional[str] = Query(None, description="Tìm theo tên, email, sđt, công ty"),
    status: Optional[str] = Query(None, description="UNASSIGNED, ASSIGNED, CONTACTED, QUALIFIED, REJECTED, CONVERTED"),
    score_grade: Optional[str] = Query(None, description="Hot, Warm, Cold"),
    source: Optional[str] = Query(None, description="Nguồn: Web Form, Excel Import, v.v."),
    campaign_id: Optional[str] = Query(None, description="Lọc theo mã chiến dịch"),
    sla_overdue: Optional[bool] = Query(None, description="Lọc các lead bị quá hạn SLA"),
    page: int = Query(1, ge=1, description="Số thứ tự trang"),
    limit: int = Query(25, ge=1, le=200, description="Số dòng mỗi trang"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> List[LeadDTO]:
    skip = (page - 1) * limit
    dtos, total = LeadService.get_leads(
        db=db,
        user=current_user,
        search=search,
        status=status,
        score_grade=score_grade,
        source=source,
        campaign_id=campaign_id,
        sla_overdue=sla_overdue,
        skip=skip,
        limit=limit,
    )
    total_pages = (total + limit - 1) // limit if total > 0 else 1

    response.headers["X-Total-Count"] = str(total)
    response.headers["X-Total-Pages"] = str(total_pages)
    response.headers["X-Current-Page"] = str(page)
    response.headers["X-Page-Size"] = str(limit)

    return dtos


@router.post("", response_model=LeadDTO, status_code=status.HTTP_201_CREATED, summary="Tạo mới Đầu mối bán hàng (S4-02)")
def create_lead(
    dto: CreateLeadDTO,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    lead = LeadService.create_lead(db=db, dto=dto, user=current_user)
    return LeadService.map_to_dto(lead)


@router.get("/{lead_id}", response_model=LeadDTO, summary="Xem chi tiết Đầu mối")
def get_lead_detail(
    lead_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    return LeadService.get_lead_by_id(db=db, lead_id=lead_id)


@router.put("/{lead_id}", response_model=LeadDTO, summary="Cập nhật thông tin Đầu mối")
def update_lead(
    lead_id: str,
    dto: UpdateLeadDTO,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> LeadDTO:
    return LeadService.update_lead(db=db, lead_id=lead_id, dto=dto)


@router.delete("/{lead_id}", status_code=status.HTTP_204_NO_CONTENT, summary="Xóa Đầu mối")
def delete_lead(
    lead_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead.")
    db.delete(lead)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)
