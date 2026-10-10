"""
Module Xuất Báo Cáo Excel Chuyên Nghiệp - Ticket SCRUM-99
Sử dụng openpyxl để tạo bảng biểu đẹp mắt, định dạng số tiền VND, tỷ lệ % và auto-fit độ rộng cột.
"""

import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from database import format_date, format_datetime


def generate_excel_report(db, from_dt=None, to_dt=None, date_label="Toàn bộ thời gian"):
    """
    Tạo file Excel gồm 3 Sheets chi tiết:
    - Sheet 1: Hiệu Quả Theo Nguồn Lead
    - Sheet 2: Hiệu Quả Theo Chiến Dịch Marketing
    - Sheet 3: Danh Sách Lead Chi Tiết
    """
    wb = openpyxl.Workbook()

    # Định nghĩa Styles dùng chung
    font_title = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=11, color="0F172A")

    fill_header = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    fill_subtotal = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_highlight = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )
    double_bottom_side = Side(border_style="double", color="0F172A")
    border_total = Border(
        top=thin_border_side,
        bottom=double_bottom_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    num_format_currency = '#,##0 "đ"'
    num_format_percent = '0.0%'
    num_format_int = '#,##0'

    # =========================================================================
    # SHEET 1: BÁO CÁO HIỆU QUẢ THEO NGUỒN LEAD
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Hiệu Quả Nguồn Lead"
    ws1.views.sheetView[0].showGridLines = True

    # Tiêu đề báo cáo
    ws1["A1"] = "BÁO CÁO HIỆU QUẢ TỪNG NGUỒN LEAD & ĐÁNH GIÁ DOANH THU"
    ws1["A1"].font = font_title
    ws1["A2"] = f"Mã Nhiệm Vụ: SCRUM-30 / SCRUM-99 | Khoảng thời gian: {date_label} | Xuất lúc: {datetime.now().strftime('%H:%M:%S %d/%m/%Y')}"
    ws1["A2"].font = font_subtitle

    # Headers Bảng 1
    headers1 = [
        "STT",
        "Nguồn Lead",
        "Số Lead",
        "Lead Sales Đã Nhận",
        "Tỷ Lệ Nhận (%)",
        "Chuyển Đổi Thành Cơ Hội",
        "Tỷ Lệ Ra Cơ Hội (%)",
        "Hợp Đồng Thành Công",
        "Doanh Thu Thực Tế (VNĐ)",
        "Chi Phí Marketing (VNĐ)",
        "ROAS (Doanh Thu/Chi Phí)",
        "Đánh Giá & Khuyến Nghị Dồn Ngân Sách"
    ]

    header_row = 4
    for col_idx, h in enumerate(headers1, start=1):
        cell = ws1.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    source_reports = db.get_report_by_source(from_dt, to_dt)
    current_row = header_row + 1

    total_leads_sum = 0
    total_accepted_sum = 0
    total_opps_sum = 0
    total_won_sum = 0
    total_rev_sum = 0
    total_spent_sum = 0

    for idx, s in enumerate(source_reports, start=1):
        ws1.cell(row=current_row, column=1, value=idx).alignment = align_center
        ws1.cell(row=current_row, column=2, value=s["source"]).alignment = align_left
        
        c3 = ws1.cell(row=current_row, column=3, value=s["total_leads"])
        c3.number_format = num_format_int
        c3.alignment = align_right

        c4 = ws1.cell(row=current_row, column=4, value=s["accepted_leads"])
        c4.number_format = num_format_int
        c4.alignment = align_right

        c5 = ws1.cell(row=current_row, column=5, value=s["acceptance_rate"] / 100)
        c5.number_format = num_format_percent
        c5.alignment = align_right

        c6 = ws1.cell(row=current_row, column=6, value=s["opportunity_leads"])
        c6.number_format = num_format_int
        c6.alignment = align_right

        c7 = ws1.cell(row=current_row, column=7, value=s["opportunity_conversion_rate"] / 100)
        c7.number_format = num_format_percent
        c7.alignment = align_right

        c8 = ws1.cell(row=current_row, column=8, value=s["won_deals"])
        c8.number_format = num_format_int
        c8.alignment = align_right

        c9 = ws1.cell(row=current_row, column=9, value=s["actual_revenue"])
        c9.number_format = num_format_currency
        c9.alignment = align_right

        c10 = ws1.cell(row=current_row, column=10, value=s["spent"])
        c10.number_format = num_format_currency
        c10.alignment = align_right

        c11 = ws1.cell(row=current_row, column=11, value=s["roas"])
        c11.number_format = '0.00 "lần"'
        c11.alignment = align_right

        c12 = ws1.cell(row=current_row, column=12, value=s["recommendation"])
        c12.alignment = align_left
        if s["recommendation_code"] == "INCREASE":
            c12.fill = fill_highlight
            c12.font = font_bold

        for col in range(1, len(headers1) + 1):
            cell = ws1.cell(row=current_row, column=col)
            cell.border = border_cell
            if col != 12 and not cell.font.bold:
                cell.font = font_regular

        total_leads_sum += s["total_leads"]
        total_accepted_sum += s["accepted_leads"]
        total_opps_sum += s["opportunity_leads"]
        total_won_sum += s["won_deals"]
        total_rev_sum += s["actual_revenue"]
        total_spent_sum += s["spent"]

        current_row += 1

    # Dòng Tổng Cộng
    ws1.cell(row=current_row, column=1, value="")
    tot_label = ws1.cell(row=current_row, column=2, value="TỔNG CỘNG")
    tot_label.font = font_bold
    tot_label.alignment = align_left

    t_leads = ws1.cell(row=current_row, column=3, value=total_leads_sum)
    t_leads.number_format = num_format_int
    t_leads.font = font_bold
    t_leads.alignment = align_right

    t_acc = ws1.cell(row=current_row, column=4, value=total_accepted_sum)
    t_acc.number_format = num_format_int
    t_acc.font = font_bold
    t_acc.alignment = align_right

    avg_acc_rate = (total_accepted_sum / total_leads_sum) if total_leads_sum > 0 else 0
    t_acc_r = ws1.cell(row=current_row, column=5, value=avg_acc_rate)
    t_acc_r.number_format = num_format_percent
    t_acc_r.font = font_bold
    t_acc_r.alignment = align_right

    t_opp = ws1.cell(row=current_row, column=6, value=total_opps_sum)
    t_opp.number_format = num_format_int
    t_opp.font = font_bold
    t_opp.alignment = align_right

    avg_opp_rate = (total_opps_sum / total_leads_sum) if total_leads_sum > 0 else 0
    t_opp_r = ws1.cell(row=current_row, column=7, value=avg_opp_rate)
    t_opp_r.number_format = num_format_percent
    t_opp_r.font = font_bold
    t_opp_r.alignment = align_right

    t_won = ws1.cell(row=current_row, column=8, value=total_won_sum)
    t_won.number_format = num_format_int
    t_won.font = font_bold
    t_won.alignment = align_right

    t_rev = ws1.cell(row=current_row, column=9, value=total_rev_sum)
    t_rev.number_format = num_format_currency
    t_rev.font = font_bold
    t_rev.alignment = align_right

    t_sp = ws1.cell(row=current_row, column=10, value=total_spent_sum)
    t_sp.number_format = num_format_currency
    t_sp.font = font_bold
    t_sp.alignment = align_right

    ov_roas = (total_rev_sum / total_spent_sum) if total_spent_sum > 0 else 0
    t_roas = ws1.cell(row=current_row, column=11, value=ov_roas)
    t_roas.number_format = '0.00 "lần"'
    t_roas.font = font_bold
    t_roas.alignment = align_right

    ws1.cell(row=current_row, column=12, value="")

    for col in range(1, len(headers1) + 1):
        cell = ws1.cell(row=current_row, column=col)
        cell.fill = fill_subtotal
        cell.border = border_total

    # =========================================================================
    # SHEET 2: BÁO CÁO THEO CHIẾN DỊCH (CAMPAIGN)
    # =========================================================================
    ws2 = wb.create_sheet(title="Hiệu Quả Chiến Dịch")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "BÁO CÁO CHI TIẾT THEO CHIẾN DỊCH MARKETING"
    ws2["A1"].font = font_title
    ws2["A2"] = f"Khoảng thời gian: {date_label} | Xuất lúc: {datetime.now().strftime('%H:%M:%S %d/%m/%Y')}"
    ws2["A2"].font = font_subtitle

    headers2 = [
        "Mã Chiến Dịch",
        "Tên Chiến Dịch",
        "Nguồn Lead",
        "Ngân Sách (VNĐ)",
        "Chi Phí Đã Dùng (VNĐ)",
        "Số Lead",
        "Tỷ Lệ Nhận (%)",
        "Số Cơ Hội",
        "Tỷ Lệ Ra Cơ Hội (%)",
        "Số Deal Won",
        "Doanh Thu Thực Tế (VNĐ)",
        "CPL (VNĐ/Lead)",
        "CPO (VNĐ/Cơ Hội)",
        "ROAS",
        "Đánh Giá"
    ]

    for col_idx, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    campaign_reports = db.get_report_by_campaign(from_dt, to_dt)
    current_row = header_row + 1

    for c in campaign_reports:
        ws2.cell(row=current_row, column=1, value=c["id"]).alignment = align_center
        ws2.cell(row=current_row, column=2, value=c["name"]).alignment = align_left
        ws2.cell(row=current_row, column=3, value=c["source"]).alignment = align_left

        c4 = ws2.cell(row=current_row, column=4, value=c["budget"])
        c4.number_format = num_format_currency
        c4.alignment = align_right

        c5 = ws2.cell(row=current_row, column=5, value=c["spent"])
        c5.number_format = num_format_currency
        c5.alignment = align_right

        c6 = ws2.cell(row=current_row, column=6, value=c["total_leads"])
        c6.number_format = num_format_int
        c6.alignment = align_right

        c7 = ws2.cell(row=current_row, column=7, value=c["acceptance_rate"] / 100)
        c7.number_format = num_format_percent
        c7.alignment = align_right

        c8 = ws2.cell(row=current_row, column=8, value=c["opportunity_leads"])
        c8.number_format = num_format_int
        c8.alignment = align_right

        c9 = ws2.cell(row=current_row, column=9, value=c["opportunity_conversion_rate"] / 100)
        c9.number_format = num_format_percent
        c9.alignment = align_right

        c10 = ws2.cell(row=current_row, column=10, value=c["won_deals"])
        c10.number_format = num_format_int
        c10.alignment = align_right

        c11 = ws2.cell(row=current_row, column=11, value=c["actual_revenue"])
        c11.number_format = num_format_currency
        c11.alignment = align_right

        c12 = ws2.cell(row=current_row, column=12, value=c["cpl"])
        c12.number_format = num_format_currency
        c12.alignment = align_right

        c13 = ws2.cell(row=current_row, column=13, value=c["cpo"])
        c13.number_format = num_format_currency
        c13.alignment = align_right

        c14 = ws2.cell(row=current_row, column=14, value=c["roas"])
        c14.number_format = '0.00 "lần"'
        c14.alignment = align_right

        c15 = ws2.cell(row=current_row, column=15, value=c["recommendation"])
        c15.alignment = align_left

        for col in range(1, len(headers2) + 1):
            cell = ws2.cell(row=current_row, column=col)
            cell.border = border_cell
            cell.font = font_regular

        current_row += 1

    # =========================================================================
    # SHEET 3: DANH SÁCH LEAD CHI TIẾT
    # =========================================================================
    ws3 = wb.create_sheet(title="Danh Sách Lead Chi Tiết")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "DANH SÁCH DỮ LIỆU LEAD CHI TIẾT TRONG KHOẢNG THỜI GIAN"
    ws3["A1"].font = font_title
    ws3["A2"] = f"Bộ lọc: {date_label} | Xuất lúc: {datetime.now().strftime('%H:%M:%S %d/%m/%Y')}"
    ws3["A2"].font = font_subtitle

    headers3 = [
        "Mã Lead",
        "Tên Công Ty / Doanh Nghiệp",
        "Người Liên Hệ",
        "Số Điện Thoại",
        "Email",
        "Nguồn Lead",
        "Chiến Dịch",
        "Ngày Tạo",
        "Trạng Thái",
        "Giá Trị Tiềm Năng (VNĐ)",
        "Doanh Thu Thực Tế (VNĐ)"
    ]

    for col_idx, h in enumerate(headers3, start=1):
        cell = ws3.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    detailed_leads = db.get_leads_filtered(from_dt, to_dt)
    current_row = header_row + 1

    for l in detailed_leads:
        ws3.cell(row=current_row, column=1, value=l["id"]).alignment = align_center
        ws3.cell(row=current_row, column=2, value=l["company_name"]).alignment = align_left
        ws3.cell(row=current_row, column=3, value=l["contact_name"]).alignment = align_left
        ws3.cell(row=current_row, column=4, value=l["phone"]).alignment = align_center
        ws3.cell(row=current_row, column=5, value=l["email"]).alignment = align_left
        ws3.cell(row=current_row, column=6, value=l["source"]).alignment = align_left
        ws3.cell(row=current_row, column=7, value=l["campaign_name"]).alignment = align_left
        ws3.cell(row=current_row, column=8, value=format_datetime(l["created_at"])).alignment = align_center
        ws3.cell(row=current_row, column=9, value=l["status"]).alignment = align_center

        c10 = ws3.cell(row=current_row, column=10, value=l["deal_value"])
        c10.number_format = num_format_currency
        c10.alignment = align_right

        c11 = ws3.cell(row=current_row, column=11, value=l["actual_revenue"])
        c11.number_format = num_format_currency
        c11.alignment = align_right

        for col in range(1, len(headers3) + 1):
            cell = ws3.cell(row=current_row, column=col)
            cell.border = border_cell
            cell.font = font_regular

        current_row += 1

    # TỰ ĐỘNG CÂN CHỈNH ĐỘ RỘNG CỘT CHO TẤT CẢ CÁC SHEET (AUTO-FIT WIDTH)
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Bỏ qua tiêu đề dài ở dòng 1 và 2 để tránh làm cột quá rộng
                if cell.row in [1, 2]:
                    continue
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Lưu workbook vào bộ nhớ đệm BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
