"""
Unit Test & Integration Test cho SCRUM-40 (Sprint 4):
Tạo Lead thủ công và Nhập Lead hàng loạt từ Excel.

Kiểm tra 100% Acceptance Criteria:
1. Tạo thủ công thành công và bắt lỗi khi thiếu source.
2. Kiểm tra định dạng số điện thoại VN và họ tên.
3. Tải file Excel mẫu chuẩn (.xlsx).
4. Xem trước (preview/dry-run), bắt lỗi dòng thiếu source, sai số điện thoại, sai email.
5. Kiểm tra trùng lặp sơ bộ (trong file và trong CSDL).
6. Hỗ trợ cả file .xlsx và .csv.
7. Thực thi import hàng loạt và rollback nếu có lỗi hệ thống.
8. Kiểm tra phân quyền vai trò (Marketing, Sales, Admin).
"""
import io
import csv
from unittest.mock import patch
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.user import User
from app.models.lead import Lead
from app.core.security import hash_password, create_access_token


@pytest.fixture
def marketing_user(db_session: Session) -> User:
    """Tạo tài khoản Nhân viên Marketing."""
    user = User(
        id="usr-mkt-scrum40",
        email="marketing_scrum40@nexuscrm.vn",
        password_hash=hash_password("MktPass123!"),
        full_name="Nguyễn Thị Mai (Marketing)",
        role="MARKETING",
        department="Phòng Marketing & Truyền Thông",
        status="active",
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def sales_user(db_session: Session) -> User:
    """Tạo tài khoản Nhân viên Kinh doanh."""
    user = User(
        id="usr-sales-scrum40",
        email="sales_scrum40@nexuscrm.vn",
        password_hash=hash_password("SalesPass123!"),
        full_name="Trần Văn Nam (Sales)",
        role="SALES_REP",
        department="Phòng Kinh Doanh",
        status="active",
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def mkt_headers(marketing_user: User) -> dict:
    token = create_access_token(marketing_user.id, marketing_user.email, marketing_user.role)
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def sales_headers(sales_user: User) -> dict:
    token = create_access_token(sales_user.id, sales_user.email, sales_user.role)
    return {"Authorization": f"Bearer {token}"}


def make_excel_file(headers: list, rows: list) -> io.BytesIO:
    """Helper tạo tệp Excel .xlsx giả lập trong bộ nhớ."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def make_csv_file(headers: list, rows: list) -> io.BytesIO:
    """Helper tạo tệp CSV giả lập trong bộ nhớ."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    buf = io.BytesIO(out.getvalue().encode("utf-8-sig"))
    buf.seek(0)
    return buf


class TestLeadManualCreation:
    """Test Suite cho tính năng Tạo Lead thủ công (Single Lead Creation)."""

    def test_create_lead_manual_success(self, client: TestClient, mkt_headers: dict, db_session: Session):
        """Tạo thủ công thành công với đầy đủ thông tin hợp lệ và có nguồn."""
        payload = {
            "full_name": "Nguyễn Hoàng Long",
            "phone": "0912345678",
            "email": "hoanglong@techcorp.vn",
            "company": "Tech Corp Vietnam",
            "interest_need": "Phần mềm CRM quản trị quan hệ khách hàng B2B",
            "source": "Hội thảo Chuyển đổi số 2026",
            "notes": "Gặp tại gian hàng A1, quan tâm triển khai trong Q2",
        }
        res = client.post("/api/v1/leads/manual", json=payload, headers=mkt_headers)
        assert res.status_code == 201
        data = res.json()
        assert data["full_name"] == "Nguyễn Hoàng Long"
        assert data["phone"] == "0912345678"
        assert data["email"] == "hoanglong@techcorp.vn"
        assert data["company"] == "Tech Corp Vietnam"
        assert data["source"] == "Hội thảo Chuyển đổi số 2026"
        assert data["status"] == "NEW"
        assert data["created_by"] == "usr-mkt-scrum40"

        # Kiểm tra bản ghi đã lưu vào CSDL
        lead_in_db = db_session.query(Lead).filter(Lead.id == data["id"]).first()
        assert lead_in_db is not None
        assert lead_in_db.source == "Hội thảo Chuyển đổi số 2026"

    def test_create_lead_manual_missing_source_fails(self, client: TestClient, mkt_headers: dict):
        """Mọi lead nhập vào BẮT BUỘC phải có nguồn - nếu để trống source hệ thống từ chối."""
        # 1. Để trống chuỗi rỗng
        payload_empty = {
            "full_name": "Trần Thị Mai",
            "phone": "0987654321",
            "source": "   ",
        }
        res_empty = client.post("/api/v1/leads/manual", json=payload_empty, headers=mkt_headers)
        assert res_empty.status_code == 422

        # 2. Không truyền trường source
        payload_none = {
            "full_name": "Trần Thị Mai",
            "phone": "0987654321",
        }
        res_none = client.post("/api/v1/leads/manual", json=payload_none, headers=mkt_headers)
        assert res_none.status_code == 422

    def test_create_lead_manual_invalid_phone_fails(self, client: TestClient, mkt_headers: dict):
        """Bắt lỗi khi số điện thoại không đúng định dạng VN (10 chữ số)."""
        payload = {
            "full_name": "Lê Văn Hùng",
            "phone": "123456",  # Quá ngắn
            "source": "Danh thiếp",
        }
        res = client.post("/api/v1/leads/manual", json=payload, headers=mkt_headers)
        assert res.status_code == 422

    def test_create_lead_manual_invalid_email_fails(self, client: TestClient, mkt_headers: dict):
        """Bắt lỗi khi email sai định dạng."""
        payload = {
            "full_name": "Phạm Quốc Bảo",
            "phone": "0905123456",
            "email": "not-an-email",
            "source": "Sự kiện",
        }
        res = client.post("/api/v1/leads/manual", json=payload, headers=mkt_headers)
        assert res.status_code == 422


class TestLeadImportTemplate:
    """Test Suite cho tính năng Tải file Excel mẫu chuẩn."""

    def test_download_excel_template(self, client: TestClient, sales_headers: dict):
        """Tải về tệp .xlsx chuẩn gồm các cột quy định và dữ liệu mẫu minh họa."""
        res = client.get("/api/v1/leads/import/template", headers=sales_headers)
        assert res.status_code == 200
        assert "spreadsheetml.sheet" in res.headers["content-type"]
        assert "lead_import_template.xlsx" in res.headers.get("content-disposition", "")

        # Mở file Excel từ buffer để kiểm tra cấu trúc
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        assert ws.title == "DanhSachLead"

        headers = [ws.cell(row=1, column=col).value for col in range(1, 8)]
        assert "Họ và tên (*)" in headers
        assert "Số điện thoại (*)" in headers
        assert "Nguồn lead (*)" in headers

        # Kiểm tra có ít nhất 1 dòng dữ liệu mẫu
        sample_name = ws.cell(row=2, column=1).value
        assert sample_name is not None and len(sample_name) > 0


class TestLeadImportPreview:
    """Test Suite cho tính năng Xem trước và Báo lỗi chi tiết từng dòng (Preview/Dry-run)."""

    def test_preview_valid_file(self, client: TestClient, mkt_headers: dict):
        """Xem trước tệp hợp lệ 100%."""
        headers = ["Họ và tên", "Số điện thoại", "Email", "Công ty", "Nhu cầu quan tâm", "Nguồn lead", "Ghi chú"]
        rows = [
            ["Vũ Minh Tuấn", "0911223344", "tuan.vu@vcorp.vn", "V-Corp", "ERP Cloud", "Hội thảo", "Gặp hội thảo"],
            ["Đỗ Lan Anh", "0933445566", "lananh@fintech.vn", "Fintech VN", "CRM Pro", "Danh thiếp", "Danh thiếp"],
        ]
        buf = make_excel_file(headers, rows)

        res = client.post(
            "/api/v1/leads/import/preview",
            files={"file": ("leads.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=mkt_headers,
        )
        assert res.status_code == 200
        data = res.json()
        assert data["total_rows"] == 2
        assert data["valid_count"] == 2
        assert data["invalid_count"] == 0
        for r in data["preview_rows"]:
            assert r["is_valid"] is True
            assert len(r["errors"]) == 0

    def test_preview_detects_missing_source_and_invalid_data(self, client: TestClient, mkt_headers: dict):
        """Báo lỗi theo từng dòng: thiếu source, sai số ĐT, thiếu tên, sai email."""
        headers = ["Họ và tên", "Số điện thoại", "Email", "Công ty", "Nhu cầu", "Nguồn", "Ghi chú"]
        rows = [
            ["Nguyễn Hợp Lệ", "0912345678", "valid@test.com", "Test Corp", "CRM", "Hội thảo", "OK"],
            ["Trần Thiếu Nguồn", "0987654321", "ok@test.com", "Test Corp", "CRM", "", "Không có nguồn"],
            ["Lê Sai Điện Thoại", "0123456", "ok2@test.com", "Test Corp", "CRM", "Danh thiếp", "SĐT sai"],
            ["", "0905112233", "ok3@test.com", "Test Corp", "CRM", "Sự kiện", "Thiếu tên"],
            ["Phạm Sai Email", "0977889900", "invalid-email-address", "Test Corp", "CRM", "Giới thiệu", "Email sai"],
        ]
        buf = make_excel_file(headers, rows)

        res = client.post(
            "/api/v1/leads/import/preview",
            files={"file": ("leads_errors.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=mkt_headers,
        )
        assert res.status_code == 200
        data = res.json()
        assert data["total_rows"] == 5
        assert data["valid_count"] == 1
        assert data["invalid_count"] == 4

        preview_rows = data["preview_rows"]
        # Dòng 1 hợp lệ
        assert preview_rows[0]["is_valid"] is True

        # Dòng 2 thiếu nguồn (source)
        assert preview_rows[1]["is_valid"] is False
        assert any("Nguồn lead (source) là bắt buộc" in err for err in preview_rows[1]["errors"])

        # Dòng 3 sai số điện thoại
        assert preview_rows[2]["is_valid"] is False
        assert any("Số điện thoại không đúng định dạng" in err for err in preview_rows[2]["errors"])

        # Dòng 4 thiếu họ tên
        assert preview_rows[3]["is_valid"] is False
        assert any("Họ và tên là bắt buộc" in err for err in preview_rows[3]["errors"])

        # Dòng 5 sai định dạng email
        assert preview_rows[4]["is_valid"] is False
        assert any("Email không đúng định dạng" in err for err in preview_rows[4]["errors"])

    def test_preview_duplicate_detection(self, client: TestClient, mkt_headers: dict, db_session: Session):
        """Phát hiện trùng lặp sơ bộ trong cùng file và với CSDL hiện có."""
        # Tạo sẵn 1 lead trong DB
        existing = Lead(
            id="lead-db-dup",
            full_name="Nguyễn Đã Có Trong DB",
            phone="0911223344",
            email="existing@db.vn",
            source="Sự kiện cũ",
            status="NEW",
        )
        db_session.add(existing)
        db_session.commit()

        headers = ["Họ và tên", "Số điện thoại", "Email", "Nguồn lead"]
        rows = [
            ["Trùng SĐT Trong DB", "0911223344", "unique1@test.com", "Hội thảo"],
            ["Trùng Email Trong DB", "0922334455", "existing@db.vn", "Hội thảo"],
            ["Dòng Gốc File", "0977889911", "filedup@test.com", "Hội thảo"],
            ["Trùng Nội Bộ File", "0977889911", "another@test.com", "Danh thiếp"],
        ]
        buf = make_excel_file(headers, rows)

        res = client.post(
            "/api/v1/leads/import/preview",
            files={"file": ("dup_check.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=mkt_headers,
        )
        assert res.status_code == 200
        data = res.json()
        assert data["valid_count"] == 1  # Chỉ dòng 3 (dòng gốc đầu tiên) là hợp lệ
        assert data["invalid_count"] == 3

        p_rows = data["preview_rows"]
        assert any("Số điện thoại đã tồn tại trên hệ thống" in err for err in p_rows[0]["errors"])
        assert any("Email đã tồn tại trên hệ thống" in err for err in p_rows[1]["errors"])
        assert p_rows[2]["is_valid"] is True
        assert any("Số điện thoại trùng lặp với dòng" in err for err in p_rows[3]["errors"])

    def test_preview_csv_support(self, client: TestClient, sales_headers: dict):
        """Xem trước tệp CSV có hỗ trợ tiếng Việt có dấu (UTF-8)."""
        headers = ["Họ và tên", "Số điện thoại", "Email", "Công ty", "Nguồn lead"]
        rows = [
            ["Nguyễn Văn CSV", "0918889999", "csv@test.vn", "CSV Tech", "Hội thảo"],
        ]
        buf = make_csv_file(headers, rows)

        res = client.post(
            "/api/v1/leads/import/preview",
            files={"file": ("leads.csv", buf.getvalue(), "text/csv")},
            headers=sales_headers,
        )
        assert res.status_code == 200
        data = res.json()
        assert data["total_rows"] == 1
        assert data["valid_count"] == 1
        assert data["preview_rows"][0]["full_name"] == "Nguyễn Văn CSV"


class TestLeadImportExecute:
    """Test Suite cho tính năng Thực thi nhập dữ liệu theo Transaction."""

    def test_execute_import_success(self, client: TestClient, sales_headers: dict, db_session: Session):
        """Nhập thành công danh sách các dòng hợp lệ, bỏ qua các dòng lỗi."""
        payload = {
            "rows": [
                {
                    "row_number": 2,
                    "full_name": "Lê Anh Tuấn",
                    "phone": "0981112222",
                    "email": "tuan.le@fpt.vn",
                    "company": "FPT Telecom",
                    "interest_need": "Cloud Server & CRM",
                    "source": "Hội thảo FPT TechDay",
                    "notes": "Quan tâm dịch vụ tháng tới",
                    "is_valid": True,
                    "errors": [],
                },
                {
                    "row_number": 3,
                    "full_name": "Phạm Thu Hà",
                    "phone": "0972223333",
                    "email": "ha.pham@viettel.vn",
                    "company": "Viettel Post",
                    "interest_need": "Quản lý khách hàng",
                    "source": "Danh thiếp",
                    "notes": None,
                    "is_valid": True,
                    "errors": [],
                },
                {
                    "row_number": 4,
                    "full_name": "Dòng Lỗi Bị Bỏ Qua",
                    "phone": "123",
                    "source": None,
                    "is_valid": False,
                    "errors": ["Nguồn lead (source) là bắt buộc"],
                },
            ],
            "skip_errors": True,
        }

        res = client.post("/api/v1/leads/import/execute", json=payload, headers=sales_headers)
        assert res.status_code == 200
        data = res.json()
        assert data["total_rows"] == 3
        assert data["imported_count"] == 2
        assert data["failed_count"] == 1

        # Xác minh trong CSDL
        tuan = db_session.query(Lead).filter(Lead.phone == "0981112222").first()
        assert tuan is not None
        assert tuan.full_name == "Lê Anh Tuấn"
        assert tuan.status == "NEW"
        assert tuan.source == "Hội thảo FPT TechDay"
        assert tuan.created_by == "usr-sales-scrum40"

        ha = db_session.query(Lead).filter(Lead.phone == "0972223333").first()
        assert ha is not None
        assert ha.source == "Danh thiếp"

    def test_execute_transaction_rollback_on_error(self, client: TestClient, sales_headers: dict, db_session: Session):
        """Kiểm tra cơ chế Transaction Rollback nếu gặp lỗi hệ thống giữa chừng."""
        payload = {
            "rows": [
                {
                    "row_number": 2,
                    "full_name": "Trần Rollback",
                    "phone": "0934567890",
                    "source": "Hội thảo",
                    "is_valid": True,
                    "errors": [],
                }
            ],
            "skip_errors": True,
        }

        # Giả lập lỗi xảy ra khi db.commit()
        with patch.object(Session, "commit", side_effect=Exception("Database connection severed")):
            res = client.post("/api/v1/leads/import/execute", json=payload, headers=sales_headers)
            assert res.status_code == 500
            assert "Lỗi transaction" in res.json()["detail"]

        # Kiểm tra không có dữ liệu nào bị lưu sót lại sau rollback
        record = db_session.query(Lead).filter(Lead.phone == "0934567890").first()
        assert record is None


class TestLeadAuthorization:
    """Test Suite cho phân quyền người dùng (Role RBAC)."""

    def test_unauthorized_role_forbidden(self, client: TestClient, db_session: Session):
        """Tài khoản có vai trò không được cấp phép (ví dụ: GUEST) bị từ chối truy cập 403."""
        guest_user = User(
            id="usr-guest-test",
            email="guest@test.vn",
            password_hash=hash_password("Pass123!"),
            full_name="Khách Vãng Lai",
            role="GUEST",
            status="active",
        )
        db_session.add(guest_user)
        db_session.commit()

        token = create_access_token(guest_user.id, guest_user.email, guest_user.role)
        headers = {"Authorization": f"Bearer {token}"}

        res = client.post(
            "/api/v1/leads/manual",
            json={"full_name": "Test User", "phone": "0912345678", "source": "Hội thảo"},
            headers=headers,
        )
        assert res.status_code == 403
        assert "Bạn không có quyền thao tác quản lý Lead" in res.json()["detail"]
