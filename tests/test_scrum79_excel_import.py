"""
Bộ kiểm thử tự động toàn diện cho SCRUM-79 / SCRUM-123 BE:
Tính năng Nhập người dùng hàng loạt từ tệp Excel (Template, Preview & Execute).
Bảo đảm 100% yêu cầu bảo mật RBAC, validation, data scope và transaction an toàn.
"""
import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.user import User
from app.core.security import hash_password, verify_password, create_access_token


@pytest.fixture
def test_admin_user(db_session: Session) -> User:
    user = User(
        id="usr-scrum79-admin",
        email="admin_scrum79@nexuscrm.vn",
        password_hash=hash_password("AdminPass123!"),
        full_name="Quản Trị Viên SCRUM-79",
        role="Super Admin",
        status="active",
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def test_staff_user(db_session: Session) -> User:
    user = User(
        id="usr-scrum79-staff",
        email="staff_scrum79@nexuscrm.vn",
        password_hash=hash_password("StaffPass123!"),
        full_name="Nhân Viên Kinh Doanh SCRUM-79",
        role="sales",
        status="active",
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def admin_headers(test_admin_user: User) -> dict:
    token = create_access_token(test_admin_user.id, test_admin_user.email, test_admin_user.role)
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def staff_headers(test_staff_user: User) -> dict:
    token = create_access_token(test_staff_user.id, test_staff_user.email, test_staff_user.role)
    return {"Authorization": f"Bearer {token}"}


def create_excel_file(headers: list, rows: list) -> io.BytesIO:
    """Helper tạo buffer file Excel .xlsx trong bộ nhớ."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ==============================================================================
# 1. TEST TẢI FILE EXCEL MẪU (DOWNLOAD TEMPLATE)
# ==============================================================================

def test_download_template_success(client: TestClient, admin_headers: dict):
    response = client.get("/api/v1/users/import/template", headers=admin_headers)
    assert response.status_code == 200
    assert "users_import_template.xlsx" in response.headers.get("content-disposition", "")
    assert "spreadsheetml" in response.headers.get("content-type", "")

    # Đọc lại nội dung file excel trả về
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "User_Import_Template"

    # Kiểm tra các cột tiêu đề
    headers = [cell.value for cell in ws[1]]
    assert headers == ["full_name", "email", "phone", "role", "department"]

    # Kiểm tra có ít nhất 2 dòng mẫu
    assert ws.max_row >= 3


def test_download_template_forbidden_for_staff(client: TestClient, staff_headers: dict):
    response = client.get("/api/v1/users/import/template", headers=staff_headers)
    assert response.status_code == 403
    assert "không có quyền" in response.json()["detail"]


# ==============================================================================
# 2. TEST XEM TRƯỚC & KIỂM TRA DỮ LIỆU (PREVIEW & VALIDATE)
# ==============================================================================

def test_preview_invalid_file_extension(client: TestClient, admin_headers: dict):
    files = {"file": ("test.txt", io.BytesIO(b"Hello world"), "text/plain")}
    response = client.post("/api/v1/users/import/preview", files=files, headers=admin_headers)
    assert response.status_code == 400
    assert "Chỉ chấp nhận tệp định dạng Excel" in response.json()["detail"]


def test_preview_excel_validation(client: TestClient, admin_headers: dict, test_admin_user: User):
    headers = ["full_name", "email", "phone", "role", "department"]
    rows = [
        # 1. Hợp lệ
        ["Vũ Hải Đăng", "haidang_s79@nexuscrm.vn", "0901234567", "Account Executive", "Kinh Doanh Miền Nam"],
        # 2. Lỗi format email
        ["Trần Văn Lỗi Email", "invalid-email-format", "0901234567", "sales", "Kinh Doanh"],
        # 3. Lỗi thiếu Họ và tên
        ["", "noname_s79@nexuscrm.vn", "0901234567", "sales", "Kinh Doanh"],
        # 4. Lỗi trùng email trong file (trùng với dòng 1)
        ["Người Trùng Trong File", "haidang_s79@nexuscrm.vn", "0901234567", "sales", "Kinh Doanh"],
        # 5. Lỗi trùng email đã có trong DB
        ["Trùng Với Admin", test_admin_user.email, "0901234567", "admin", "Ban Quản Trị"],
        # 6. Lỗi số điện thoại không chuẩn
        ["Lê Lỗi Phone", "phone_err_s79@nexuscrm.vn", "12345", "sales", "Kinh Doanh"],
        # 7. Lỗi vai trò không hợp lệ
        ["Phạm Lỗi Role", "role_err_s79@nexuscrm.vn", "0901234567", "UnknownRole999", "Kinh Doanh"],
    ]

    excel_file = create_excel_file(headers, rows)
    files = {"file": ("users_test.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    response = client.post("/api/v1/users/import/preview", files=files, headers=admin_headers)
    assert response.status_code == 200
    data = response.json()

    assert data["total_rows"] == 7
    assert data["valid_count"] == 1
    assert data["error_count"] == 6

    details = data["details"]
    assert len(details) == 7

    # Dòng 1: VALID
    assert details[0]["status"] == "VALID"
    assert len(details[0]["errors"]) == 0
    assert details[0]["email"] == "haidang_s79@nexuscrm.vn"

    # Dòng 2: INVALID (Email)
    assert details[1]["status"] == "INVALID"
    assert any("không đúng định dạng" in err for err in details[1]["errors"])

    # Dòng 3: INVALID (Họ tên)
    assert details[2]["status"] == "INVALID"
    assert any("Họ và tên không được để trống" in err for err in details[2]["errors"])

    # Dòng 4: INVALID (Trùng file)
    assert details[3]["status"] == "INVALID"
    assert any("trùng lặp trong tệp Excel" in err for err in details[3]["errors"])

    # Dòng 5: INVALID (Trùng DB)
    assert details[4]["status"] == "INVALID"
    assert any("đã tồn tại trên hệ thống" in err for err in details[4]["errors"])

    # Dòng 6: INVALID (Phone)
    assert details[5]["status"] == "INVALID"
    assert any("Số điện thoại" in err for err in details[5]["errors"])

    # Dòng 7: INVALID (Role)
    assert details[6]["status"] == "INVALID"
    assert any("không hợp lệ" in err for err in details[6]["errors"])


# ==============================================================================
# 3. TEST THỰC THI NHẬP DỮ LIỆU (EXECUTE IMPORT)
# ==============================================================================

def test_execute_import_from_excel_file(client: TestClient, admin_headers: dict, db_session: Session):
    headers = ["full_name", "email", "phone", "role", "department"]
    rows = [
        # Dòng 1: Hợp lệ -> Được nhập
        ["Nguyễn Văn Thực Thi", "thucthi_ok@nexuscrm.vn", "0912345678", "Sales Manager", "Miền Bắc (Hà Nội)"],
        # Dòng 2: Lỗi Email -> Bị bỏ qua
        ["Trần Lỗi", "bademail", "0912345678", "sales", "Miền Bắc"],
    ]

    excel_file = create_excel_file(headers, rows)
    files = {"file": ("import_exec.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    response = client.post("/api/v1/users/import/execute", files=files, headers=admin_headers)
    assert response.status_code == 200
    data = response.json()

    assert data["total_rows"] == 2
    assert data["imported_count"] == 1
    assert data["failed_count"] == 1
    assert len(data["imported_users"]) == 1
    assert len(data["failed_rows"]) == 1

    # Kiểm tra user hợp lệ đã lưu trong DB
    user_db = db_session.query(User).filter(User.email == "thucthi_ok@nexuscrm.vn").first()
    assert user_db is not None
    assert user_db.full_name == "Nguyễn Văn Thực Thi"
    assert user_db.role == "Sales Manager"
    assert user_db.department == "Miền Bắc (Hà Nội)"
    assert user_db.data_scope == "TEAM"
    assert user_db.status == "active"
    # Kiểm tra mật khẩu mặc định được hash bằng bcrypt
    assert verify_password("Password123!", user_db.password_hash) is True

    # Kiểm tra user không hợp lệ KHÔNG được lưu trong DB
    bad_db = db_session.query(User).filter(User.email == "bademail").first()
    assert bad_db is None


def test_execute_import_from_json_payload(client: TestClient, admin_headers: dict, db_session: Session):
    payload = {
        "rows": [
            {
                "full_name": "Lê JSON Account Exec",
                "email": "json_user_s79@nexuscrm.vn",
                "role": "Account Executive",
                "department": "Khối Doanh Nghiệp",
                "phone": "0988776655",
            },
            {
                "full_name": "Trần Lỗi Phone JSON",
                "email": "json_bad_phone@nexuscrm.vn",
                "role": "Account Executive",
                "department": "Khối Doanh Nghiệp",
                "phone": "000000",
            },
        ]
    }

    response = client.post("/api/v1/users/import/execute", json=payload, headers=admin_headers)
    assert response.status_code == 200
    data = response.json()

    assert data["total_rows"] == 2
    assert data["imported_count"] == 1
    assert data["failed_count"] == 1

    user_db = db_session.query(User).filter(User.email == "json_user_s79@nexuscrm.vn").first()
    assert user_db is not None
    assert user_db.data_scope == "OWN"
    assert verify_password("Password123!", user_db.password_hash) is True
