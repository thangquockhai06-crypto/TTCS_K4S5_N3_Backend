"""
Bộ kiểm thử tự động toàn diện cho SCRUM-64 (S3-06 / Sprint 3):
Tính năng Nhập danh sách khách hàng hàng loạt từ Excel (Bulk Customer Import).
Kiểm thử 100% Acceptance Criteria:
1. GET /api/v1/customers/import/template: Tải template chuẩn .xlsx
2. POST /api/v1/customers/import/preview:
   - Validate từng dòng (name, tax_code, email, phone)
   - Quét trùng lặp CSDL (tax_code, website) và trùng nội bộ file
   - Hỗ trợ cả file .xlsx và .csv
   - Chặn tệp quá kích thước 5MB hoặc sai định dạng
3. POST /api/v1/customers/import/execute:
   - Hành động "SKIP": Bỏ qua bản ghi trùng, nạp bản ghi mới
   - Hành động "UPDATE": Cập nhật đè dữ liệu mới vào khách hàng đã có
   - Tự động bỏ qua các dòng không hợp lệ (is_valid == False)
   - Gán quyền sở hữu assigned_user_id cho người import
   - Báo cáo tổng kết CustomerImportSummaryResponse
"""
import io
import csv
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.user import User
from app.models.customer import Customer
from app.core.security import hash_password, create_access_token


@pytest.fixture
def test_sales_user(db_session: Session) -> User:
    """Tạo người dùng Sales Rep để kiểm thử."""
    user = User(
        id="usr-scrum64-sales",
        email="sales_scrum64@nexuscrm.vn",
        password_hash=hash_password("SalesPass123!"),
        full_name="Nguyễn Văn Sales SCRUM-64",
        role="sales",
        status="active",
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def sales_headers(test_sales_user: User) -> dict:
    """Tạo Bearer token cho Sales Rep."""
    token = create_access_token(test_sales_user.id, test_sales_user.email, test_sales_user.role)
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def existing_customer(db_session: Session, test_sales_user: User) -> Customer:
    """Tạo sẵn khách hàng trong CSDL để kiểm thử đối soát trùng lặp."""
    customer = Customer(
        id="cust-scrum64-exist",
        full_name="Công ty Cổ phần Công nghệ Tiên Phong",
        company="Công ty Cổ phần Công nghệ Tiên Phong",
        email="contact@tienphong.vn",
        phone="0912345678",
        tax_code="0109998888",
        website="https://tienphong.vn",
        address="123 Phố Huế, Hà Nội",
        status="lead",
        health_score=85,
        assigned_user_id=test_sales_user.id,
        is_deleted=False,
    )
    db_session.add(customer)
    db_session.commit()
    db_session.refresh(customer)
    return customer


def make_excel_buffer(headers: list, rows: list) -> io.BytesIO:
    """Helper tạo file Excel .xlsx trong bộ nhớ."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def make_csv_buffer(headers: list, rows: list) -> io.BytesIO:
    """Helper tạo file CSV trong bộ nhớ."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    buf = io.BytesIO(out.getvalue().encode("utf-8-sig"))
    buf.seek(0)
    return buf


# ============================================================================
# 1. TEST TẢI FILE EXCEL MẪU (GET /template)
# ============================================================================

def test_download_template_unauthorized(client: TestClient):
    """Không có Bearer token -> HTTP 401."""
    res = client.get("/api/v1/customers/import/template")
    assert res.status_code == 401


def test_download_template_success(client: TestClient, sales_headers: dict):
    """Tải template thành công, đúng Content-Type và cấu trúc cột."""
    res = client.get("/api/v1/customers/import/template", headers=sales_headers)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "customer_import_template.xlsx" in res.headers["content-disposition"]

    # Đọc lại nội dung Excel
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    assert ws.title == "Customer_Import_Template"

    # Kiểm tra header
    headers = [cell.value for cell in ws[1]]
    expected_headers = ["name", "tax_code", "email", "phone", "website", "address", "status"]
    assert headers == expected_headers

    # Có ít nhất 2 dòng mẫu
    assert ws.max_row >= 3


# ============================================================================
# 2. TEST XEM TRƯỚC & VALIDATE DỮ LIỆU (POST /preview)
# ============================================================================

def test_preview_valid_and_invalid_rows(client: TestClient, sales_headers: dict):
    """Kiểm tra báo lỗi chi tiết theo từng dòng (tên trống, mst sai, email sai, sđt sai)."""
    headers = ["name", "tax_code", "email", "phone", "website", "address", "status"]
    rows = [
        # Dòng 1: Hợp lệ hoàn toàn
        ["Công ty TNHH Ánh Dương", "0101234567", "info@anhduong.vn", "0901234567", "https://anhduong.vn", "Hà Nội", "CUSTOMER"],
        # Dòng 2: Thiếu tên khách hàng (bắt buộc)
        ["", "0301234567", "info@no-name.vn", "0912345678", "https://noname.vn", "Đà Nẵng", "LEAD"],
        # Dòng 3: MST sai định dạng (chỉ có 5 chữ số)
        ["Doanh nghiệp B", "12345", "test@b.vn", "0987654321", "https://b.vn", "TP.HCM", "CONTACTED"],
        # Dòng 4: Email sai định dạng
        ["Doanh nghiệp C", "0309998888", "invalid-email-format", "0934567890", "https://c.vn", "Cần Thơ", "LEAD"],
        # Dòng 5: SĐT sai định dạng
        ["Doanh nghiệp D", "0301112223", "test@d.vn", "12345", "https://d.vn", "Hải Phòng", "LEAD"],
    ]

    excel_buf = make_excel_buffer(headers, rows)
    files = {"file": ("customers.xlsx", excel_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    res = client.post("/api/v1/customers/import/preview", files=files, headers=sales_headers)
    assert res.status_code == 200
    data = res.json()

    assert data["total_rows"] == 5
    assert data["valid_rows_count"] == 1
    assert data["invalid_rows_count"] == 4

    rows_result = data["rows"]
    # Dòng 1 hợp lệ
    assert rows_result[0]["is_valid"] is True
    assert len(rows_result[0]["errors"]) == 0

    # Dòng 2 lỗi tên
    assert rows_result[1]["is_valid"] is False
    assert any("Tên khách hàng là bắt buộc" in e for e in rows_result[1]["errors"])

    # Dòng 3 lỗi tax_code
    assert rows_result[2]["is_valid"] is False
    assert any("Mã số thuế không đúng định dạng" in e for e in rows_result[2]["errors"])

    # Dòng 4 lỗi email
    assert rows_result[3]["is_valid"] is False
    assert any("Email không đúng định dạng" in e for e in rows_result[3]["errors"])

    # Dòng 5 lỗi phone
    assert rows_result[4]["is_valid"] is False
    assert any("Số điện thoại không đúng định dạng" in e for e in rows_result[4]["errors"])


def test_preview_duplicate_detection(client: TestClient, sales_headers: dict, existing_customer: Customer):
    """Kiểm tra phát hiện trùng lặp CSDL (tax_code, website) và trùng nội bộ trong file."""
    headers = ["name", "tax_code", "email", "phone", "website", "address", "status"]
    rows = [
        # Dòng 1: Trùng tax_code với existing_customer trong CSDL
        ["Công ty Tiên Phong Cập Nhật", "0109998888", "new_email@tienphong.vn", "0909999888", "https://newsite.vn", "Hà Nội", "CUSTOMER"],
        # Dòng 2: Trùng website với existing_customer trong CSDL
        ["Tiên Phong Chi Nhánh 2", "0319998888", "branch@tienphong.vn", "0901112223", "https://tienphong.vn", "TP.HCM", "LEAD"],
        # Dòng 3: Khách hàng mới không trùng
        ["Khách Hàng Mới Tinh", "0401234567", "moi@corp.vn", "0902223334", "https://moicorp.vn", "Đà Nẵng", "LEAD"],
        # Dòng 4: Trùng tax_code với Dòng 3 ngay trong cùng file
        ["Khách Hàng Trùng Dòng 3", "0401234567", "khac@corp.vn", "0905556667", "https://khac.vn", "Huế", "LEAD"],
    ]

    excel_buf = make_excel_buffer(headers, rows)
    files = {"file": ("dup_test.xlsx", excel_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    res = client.post("/api/v1/customers/import/preview", files=files, headers=sales_headers)
    assert res.status_code == 200
    data = res.json()

    assert data["total_rows"] == 4
    assert data["duplicate_rows_count"] == 3

    r1 = data["rows"][0]
    assert r1["is_duplicate"] is True
    assert r1["matched_customer_id"] == existing_customer.id
    assert "Trùng mã số thuế" in r1["duplicate_reason"]

    r2 = data["rows"][1]
    assert r2["is_duplicate"] is True
    assert r2["matched_customer_id"] == existing_customer.id
    assert "Trùng địa chỉ website" in r2["duplicate_reason"]

    r3 = data["rows"][2]
    assert r3["is_duplicate"] is False

    r4 = data["rows"][3]
    assert r4["is_duplicate"] is True
    assert "Trùng mã số thuế với dòng 4" in r4["duplicate_reason"] or "trong cùng tệp" in r4["duplicate_reason"]


def test_preview_csv_support(client: TestClient, sales_headers: dict):
    """Kiểm tra preview đọc file định dạng CSV thành công."""
    headers = ["name", "tax_code", "email", "phone", "website", "address", "status"]
    rows = [
        ["Tập đoàn Hoàng Long CSV", "0108889999", "hl@hoanglong.vn", "0908889999", "https://hoanglong.vn", "Hải Phòng", "LEAD"],
    ]
    csv_buf = make_csv_buffer(headers, rows)
    files = {"file": ("customers.csv", csv_buf, "text/csv")}

    res = client.post("/api/v1/customers/import/preview", files=files, headers=sales_headers)
    assert res.status_code == 200
    data = res.json()
    assert data["total_rows"] == 1
    assert data["valid_rows_count"] == 1
    assert data["rows"][0]["name"] == "Tập đoàn Hoàng Long CSV"


def test_preview_file_validation_errors(client: TestClient, sales_headers: dict):
    """Kiểm tra từ chối tệp không hợp lệ: tệp rỗng, sai định dạng, quá 5MB."""
    # 1. Tệp rỗng
    empty_buf = io.BytesIO(b"")
    files = {"file": ("empty.xlsx", empty_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post("/api/v1/customers/import/preview", files=files, headers=sales_headers)
    assert res.status_code == 400

    # 2. Sai phần mở rộng (.pdf)
    pdf_buf = io.BytesIO(b"PDF content")
    files = {"file": ("test.pdf", pdf_buf, "application/pdf")}
    res = client.post("/api/v1/customers/import/preview", files=files, headers=sales_headers)
    assert res.status_code == 400
    assert "Chỉ chấp nhận tệp định dạng Excel" in res.json()["detail"]

    # 3. Kích thước > 5MB
    large_buf = io.BytesIO(b"0" * (5 * 1024 * 1024 + 10))
    files = {"file": ("large.xlsx", large_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post("/api/v1/customers/import/preview", files=files, headers=sales_headers)
    assert res.status_code == 413


# ============================================================================
# 3. TEST THỰC THI NHẬP DỮ LIỆU (POST /execute)
# ============================================================================

def test_execute_import_with_skip_duplicate(client: TestClient, sales_headers: dict, existing_customer: Customer, db_session: Session):
    """
    Thực thi nhập với duplicate_action == 'SKIP':
    - Bản ghi hợp lệ mới -> được tạo mới.
    - Bản ghi trùng lặp -> bị bỏ qua, không ghi đè dữ liệu cũ.
    - Bản ghi lỗi validation -> bị failed_count.
    """
    payload = {
        "duplicate_action": "SKIP",
        "rows": [
            # Hợp lệ mới
            {
                "row_number": 2,
                "name": "Công ty TNHH Sao Mai",
                "tax_code": "0107778888",
                "email": "saomai@corp.vn",
                "phone": "0907778888",
                "website": "https://saomai.vn",
                "address": "Bắc Ninh",
                "status": "CUSTOMER",
                "is_valid": True,
                "errors": [],
                "is_duplicate": False,
            },
            # Trùng lặp với existing_customer
            {
                "row_number": 3,
                "name": "Tên Bị Ghi Đè Nếu Update",
                "tax_code": existing_customer.tax_code,
                "email": "hacked@tienphong.vn",
                "phone": "0999999999",
                "website": existing_customer.website,
                "address": "Địa chỉ mới",
                "status": "CUSTOMER",
                "is_valid": True,
                "errors": [],
                "is_duplicate": True,
                "matched_customer_id": existing_customer.id,
            },
            # Dòng lỗi validation
            {
                "row_number": 4,
                "name": "",
                "tax_code": "123",
                "email": "invalid",
                "phone": "123",
                "website": None,
                "address": None,
                "status": "LEAD",
                "is_valid": False,
                "errors": ["Tên khách hàng là bắt buộc", "Mã số thuế không đúng định dạng"],
                "is_duplicate": False,
            },
        ],
    }

    res = client.post("/api/v1/customers/import/execute", json=payload, headers=sales_headers)
    assert res.status_code == 200
    summary = res.json()

    assert summary["total_rows"] == 3
    assert summary["imported_count"] == 1
    assert summary["skipped_count"] == 1
    assert summary["failed_count"] == 1
    assert summary["updated_count"] == 0
    assert len(summary["errors_detail"]) == 1

    # Kiểm tra database: Công ty Sao Mai được tạo
    created_cust = db_session.query(Customer).filter(Customer.tax_code == "0107778888").first()
    assert created_cust is not None
    assert created_cust.full_name == "Công ty TNHH Sao Mai"
    assert created_cust.status == "active"
    assert created_cust.assigned_user_id == "usr-scrum64-sales"

    # Kiểm tra database: existing_customer KHÔNG bị thay đổi
    db_session.refresh(existing_customer)
    assert existing_customer.full_name == "Công ty Cổ phần Công nghệ Tiên Phong"
    assert existing_customer.email == "contact@tienphong.vn"


def test_execute_import_with_update_duplicate(client: TestClient, sales_headers: dict, existing_customer: Customer, db_session: Session):
    """
    Thực thi nhập với duplicate_action == 'UPDATE':
    - Bản ghi trùng lặp được cập nhật đè thông tin mới.
    """
    payload = {
        "duplicate_action": "UPDATE",
        "rows": [
            {
                "row_number": 2,
                "name": "Công ty CP Công nghệ Tiên Phong (Đã đổi tên)",
                "tax_code": existing_customer.tax_code,
                "email": "new_updated@tienphong.vn",
                "phone": "0988776655",
                "website": "https://tienphong-new.vn",
                "address": "Tòa nhà Keangnam, Hà Nội",
                "status": "CUSTOMER",
                "is_valid": True,
                "errors": [],
                "is_duplicate": True,
                "matched_customer_id": existing_customer.id,
            }
        ],
    }

    res = client.post("/api/v1/customers/import/execute", json=payload, headers=sales_headers)
    assert res.status_code == 200
    summary = res.json()

    assert summary["total_rows"] == 1
    assert summary["updated_count"] == 1
    assert summary["imported_count"] == 0
    assert summary["skipped_count"] == 0
    assert summary["failed_count"] == 0

    # Kiểm tra database: existing_customer đã được cập nhật
    db_session.refresh(existing_customer)
    assert existing_customer.full_name == "Công ty CP Công nghệ Tiên Phong (Đã đổi tên)"
    assert existing_customer.email == "new_updated@tienphong.vn"
    assert existing_customer.phone == "0988776655"
    assert existing_customer.website == "https://tienphong-new.vn"
    assert existing_customer.address == "Tòa nhà Keangnam, Hà Nội"
    assert existing_customer.status == "active"
